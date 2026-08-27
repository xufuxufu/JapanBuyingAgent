from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.db import build_engine
from app.local_product import (
    derive_jan_from_qinsi_sku,
    ensure_qinsi_derived_barcode,
    is_valid_jan,
    resolve_local_product_by_jan,
)
from app.models import (
    Location,
    Product,
    ProductBarcode,
    QinsiGoodsImportRow,
    QinsiImportBatch,
    QinsiInventorySnapshot,
    QinsiInventorySnapshotLine,
    QinsiMasterValue,
)
from app.product_identity import format_product_display_name, normalize_product_name_whitespace
from app.product_image_localization import product_image_summary, queue_product_image_localization
from app.product_merge import merge_duplicate_jan_group, merge_products_with_jan_into_primary


MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
PARSE_VERSION = 3
PRODUCT_SHEET = "商品导入"
CONFIG_SHEET = "配置"

FIXED_HEADERS = (
    "名称（必填）", "商品规格", "货号（必填且唯一）", "条码", "型号规格", "品牌", "分类", "单位",
    "采购价", "销售价", "最低销售价", "排序", "状态", "启用积分", "库存预警下限", "库存预警上限",
    "保质期（天）", "启用批次", "过期预警（天）", "商品图片链接", "商品备注", "产地", "适用年龄",
    "商品重量（KG）", "启用序列号", "库位", "盘点库存数量", "当前库存（导入时不需要录入）",
    "盘点仓库:",
)
HEADER_ALIASES = {
    "名称(必填)": "名称（必填）",
    "货号(必填且唯一)": "货号（必填且唯一）",
    "保质期(天)": "保质期（天）",
    "过期预警(天)": "过期预警（天）",
    "商品重量(KG)": "商品重量（KG）",
    "当前库存(导入时不需要录入)": "当前库存（导入时不需要录入）",
}
HEADER_TO_FIELD = {
    "名称（必填）": "name_cn",
    "商品规格": "specification",
    "货号（必填且唯一）": "qinsi_product_code",
    "条码": "jan",
    "型号规格": "model_spec",
    "品牌": "brand",
    "分类": "category",
    "单位": "unit_name",
    "采购价": "purchase_price",
    "销售价": "sale_price",
    "最低销售价": "minimum_sale_price",
    "排序": "qinsi_sort_order",
    "状态": "status",
    "启用积分": "qinsi_points_enabled",
    "库存预警下限": "inventory_warning_lower",
    "库存预警上限": "inventory_warning_upper",
    "保质期（天）": "shelf_life_days",
    "启用批次": "batch_enabled",
    "过期预警（天）": "expiration_warning_days",
    "商品图片链接": "image_url",
    # 当前秦丝文件把商品日文名放在“商品备注”；入库后统一写 name_ja。
    "商品备注": "name_ja",
    "产地": "origin_place",
    "适用年龄": "applicable_age",
    "商品重量（KG）": "weight_kg",
    "启用序列号": "serial_number_enabled",
    "库位": "location_code",
    "盘点库存数量": "counted_inventory_quantity",
    "当前库存（导入时不需要录入）": "current_inventory_quantity",
}
FORMAL_PRODUCT_LIST_HEADERS = {
    "商品名称", "商品规格", "货号", "商品条码", "单品条码", "型号规格",
    "图片", "图片链接", "品牌", "分类", "单位", "采购价", "销售价", "最低销售价", "状态", "备注",
}
FORMAL_REQUIRED_HEADERS = {"商品名称", "货号"}
FORMAL_HEADER_TO_TEMPLATE_HEADER = {
    "商品名称": "名称（必填）",
    "商品规格": "商品规格",
    "货号": "货号（必填且唯一）",
    "型号规格": "型号规格",
    "图片": "商品图片链接",
    "图片链接": "商品图片链接",
    "品牌": "品牌",
    "分类": "分类",
    "单位": "单位",
    "采购价": "采购价",
    "销售价": "销售价",
    "最低销售价": "最低销售价",
    "状态": "状态",
    "备注": "商品备注",
}
PRODUCT_FIELDS = tuple(
    field for field in dict.fromkeys(HEADER_TO_FIELD.values())
    if field not in {"counted_inventory_quantity", "current_inventory_quantity"}
)
CONFIG_ROWS = {
    "brand": 1,
    "category": 2,
    "unit": 3,
    "warehouse": 4,
    "product_status": 5,
    "points_status": 7,
}
MASTER_PRODUCT_FIELDS = {
    "brand": ("brand", "qinsi_brand_master_id"),
    "category": ("category", "qinsi_category_master_id"),
    "unit": ("unit_name", "qinsi_unit_master_id"),
}


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    rows: list[tuple[int, dict[str, str | None]]]
    headers: tuple[str, ...]
    config: dict[str, list[str]]
    inventory_warehouse: str | None
    defined_names: tuple[str, ...]
    missing_formula_cache: dict[int, list[str]]
    sheet_name: str = PRODUCT_SHEET
    source_format: str = "qinsi_goods_template"


def _validate_archive(content: bytes) -> None:
    if not content or len(content) > MAX_XLSX_BYTES:
        raise ValueError("Excel文件为空或超过20MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES
            ):
                raise ValueError("Excel解压内容过大")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("文件不是有效的 .xlsx 工作簿")
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 .xlsx 工作簿") from exc


def _plain_decimal(value: int | float | Decimal) -> str:
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Excel数值无效：{value}") from exc
    if not number.is_finite():
        raise ValueError(f"Excel数值无效：{value}")
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _zero_format_width(number_format: str) -> int | None:
    first = (number_format or "").split(";", 1)[0]
    first = re.sub(r'"[^"]*"|\\.|_.|\*.', "", first)
    return len(first) if re.fullmatch(r"0+", first) else None


def _logical_cell_text(cell: Cell, cached_cell: Cell, *, identifier: bool = False) -> str | None:
    value = cached_cell.value
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        text = _plain_decimal(value)
        if identifier:
            width = _zero_format_width(cell.number_format)
            if width and re.fullmatch(r"\d+", text):
                text = text.zfill(width)
        return text
    text = str(value).strip()
    return text or None


def _normalized_header(value: str | None) -> str | None:
    if value is None:
        return None
    text = normalize_product_name_whitespace(value)
    return HEADER_ALIASES.get(text, text)


def _sheet_headers(values_sheet, formulas_sheet) -> list[str]:
    headers: list[str] = []
    for column in range(1, values_sheet.max_column + 1):
        header = _logical_cell_text(
            formulas_sheet.cell(1, column), values_sheet.cell(1, column),
        )
        headers.append(_normalized_header(header) or "")
    return headers


def _old_template_inventory_warehouse(headers: list[str]) -> str | None:
    canonical_fixed = tuple(_normalized_header(value) for value in headers[:-1])
    legacy_headers = tuple(header for header in FIXED_HEADERS if header != "商品规格")
    if canonical_fixed in {FIXED_HEADERS, legacy_headers}:
        return headers[-1].strip() or None
    return None


def _formal_header_score(headers: list[str]) -> int:
    header_set = {header for header in headers if header}
    if not FORMAL_REQUIRED_HEADERS <= header_set:
        return 0
    return len(header_set & FORMAL_PRODUCT_LIST_HEADERS)


def _classify_product_sheet(values_book, formulas_book) -> tuple[str, str, list[str], str | None]:
    candidates: list[tuple[int, int, str, str, list[str], str | None]] = []
    for order, sheet_name in enumerate(values_book.sheetnames):
        if sheet_name == CONFIG_SHEET:
            continue
        headers = _sheet_headers(values_book[sheet_name], formulas_book[sheet_name])
        inventory_warehouse = _old_template_inventory_warehouse(headers)
        if inventory_warehouse is not None:
            priority = 0 if sheet_name == PRODUCT_SHEET else 1
            candidates.append((priority, order, sheet_name, "qinsi_goods_template", headers, inventory_warehouse))
            continue
        score = _formal_header_score(headers)
        if score >= 5:
            priority = 2 if sheet_name == PRODUCT_SHEET else 3
            candidates.append((priority, -score, sheet_name, "qinsi_product_list", headers, None))
    if not candidates:
        raise ValueError("Excel缺少可识别的秦丝商品列表工作表")
    _priority, _order_or_score, sheet_name, source_format, headers, inventory_warehouse = sorted(candidates)[0]
    return sheet_name, source_format, headers, inventory_warehouse


def _determine_formal_jan(raw: dict[str, str | None]) -> str | None:
    for value in (raw.get("单品条码"), raw.get("商品条码"), raw.get("货号")):
        candidate = (value or "").strip()
        if is_valid_jan(candidate):
            return candidate
    return None


def _formal_row_to_template_raw(raw: dict[str, str | None]) -> dict[str, str | None]:
    mapped: dict[str, str | None] = {}
    for formal_header, template_header in FORMAL_HEADER_TO_TEMPLATE_HEADER.items():
        if formal_header in raw:
            mapped[template_header] = raw.get(formal_header)
    mapped["条码"] = _determine_formal_jan(raw)
    return mapped


def parse_qinsi_workbook(content: bytes) -> ParsedWorkbook:
    _validate_archive(content)
    try:
        values_book = load_workbook(io.BytesIO(content), data_only=True, read_only=False, keep_links=False)
        formulas_book = load_workbook(io.BytesIO(content), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Excel工作簿无法解析") from exc
    try:
        sheet_name, source_format, headers, inventory_warehouse = _classify_product_sheet(values_book, formulas_book)
        values_sheet = values_book[sheet_name]
        formulas_sheet = formulas_book[sheet_name]

        rows: list[tuple[int, dict[str, str | None]]] = []
        formula_cache_missing: dict[int, list[str]] = {}
        for row_number in range(2, values_sheet.max_row + 1):
            raw: dict[str, str | None] = {}
            missing: list[str] = []
            for column, original_header in enumerate(headers, 1):
                formula_cell = formulas_sheet.cell(row_number, column)
                cached_cell = values_sheet.cell(row_number, column)
                canonical_header = _normalized_header(original_header)
                value = _logical_cell_text(
                    formula_cell,
                    cached_cell,
                    identifier=canonical_header in {"货号（必填且唯一）", "条码"},
                )
                raw[original_header] = value
                if formula_cell.data_type == "f" and cached_cell.value is None:
                    missing.append(original_header)
            if missing:
                formula_cache_missing[row_number] = missing
            if source_format == "qinsi_product_list":
                raw = _formal_row_to_template_raw(raw)
            rows.append((row_number, raw))

        config: dict[str, list[str]] = {}
        if CONFIG_SHEET in values_book.sheetnames:
            config_values = values_book[CONFIG_SHEET]
            config_formulas = formulas_book[CONFIG_SHEET]
            for master_type, row_number in CONFIG_ROWS.items():
                values: list[str] = []
                for column in range(1, config_values.max_column + 1):
                    value = _logical_cell_text(
                        config_formulas.cell(row_number, column),
                        config_values.cell(row_number, column),
                    )
                    if value is not None:
                        values.append(value)
                config[master_type] = list(dict.fromkeys(values))
        return ParsedWorkbook(
            rows=rows,
            headers=tuple(headers),
            config=config,
            inventory_warehouse=inventory_warehouse,
            defined_names=tuple(values_book.defined_names),
            missing_formula_cache=formula_cache_missing,
            sheet_name=sheet_name,
            source_format=source_format,
        )
    finally:
        values_book.close()
        formulas_book.close()


def read_product_sheet(content: bytes) -> list[tuple[int, dict[str, str | None]]]:
    return parse_qinsi_workbook(content).rows


def _identifier(value: str | None, label: str) -> str | None:
    value = (value or "").strip()
    if not value:
        return None
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)[Ee][+-]?\d+", value):
        raise ValueError(f"{label}使用了科学计数法，无法保证原始标识符")
    return value


def _decimal_value(value: str | None, label: str, *, scale: int) -> Decimal | None:
    value = (value or "").strip()
    if not value:
        return None
    try:
        number = Decimal(value.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"{label}不是有效数值") from exc
    if not number.is_finite():
        raise ValueError(f"{label}不是有效数值")
    quantum = Decimal(1).scaleb(-scale)
    rounded = number.quantize(quantum)
    if rounded != number:
        raise ValueError(f"{label}最多支持{scale}位小数")
    return rounded


def _integer_value(value: str | None, label: str) -> int | None:
    number = _decimal_value(value, label, scale=6)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise ValueError(f"{label}必须是整数")
    return int(number)


def _boolean_value(value: str | None, label: str) -> bool | None:
    text = (value or "").strip()
    if not text:
        return None
    folded = text.casefold()
    if folded in {"启用", "是", "开启", "true", "1", "yes"}:
        return True
    if folded in {"停用", "否", "关闭", "false", "0", "no"}:
        return False
    raise ValueError(f"{label}值无效：{text}")


def _product_status(value: str | None) -> str | None:
    text = (value or "").strip()
    if not text:
        return None
    if text in {"启用", "在售", "销售中"}:
        return "active"
    if text in {"停用", "停售", "停止销售"}:
        return "inactive"
    raise ValueError(f"状态值无效：{text}")


def _canonical_raw(raw: dict[str, str | None]) -> dict[str, str | None]:
    return {
        _normalized_header(header) or header: value
        for header, value in raw.items()
    }


def _mapped_row(
    raw: dict[str, str | None],
    config: dict[str, list[str]] | None = None,
    formula_cache_missing: list[str] | None = None,
) -> tuple[dict, list[str], list[str]]:
    canonical = _canonical_raw(raw)
    mapped: dict = {}
    warnings: list[str] = []
    errors: list[str] = []
    for header, field in HEADER_TO_FIELD.items():
        value = canonical.get(header)
        try:
            if field in {"jan", "qinsi_product_code"}:
                mapped[field] = _identifier(value, "条码" if field == "jan" else "货号")
                if field == "jan" and mapped[field] and not is_valid_jan(mapped[field]):
                    errors.append("条码不是合法 JAN-8/JAN-13")
            elif field in {"purchase_price", "sale_price", "minimum_sale_price"}:
                mapped[field] = _decimal_value(value, header, scale=2)
            elif field in {"inventory_warning_lower", "inventory_warning_upper", "weight_kg"}:
                mapped[field] = _decimal_value(value, header, scale=3)
            elif field in {
                "qinsi_sort_order", "shelf_life_days", "expiration_warning_days",
                "counted_inventory_quantity", "current_inventory_quantity",
            }:
                mapped[field] = _integer_value(value, header)
            elif field in {"qinsi_points_enabled", "batch_enabled", "serial_number_enabled"}:
                mapped[field] = _boolean_value(value, header)
            elif field == "status":
                mapped[field] = _product_status(value)
            elif field in {"name_cn", "name_ja"}:
                mapped[field] = normalize_product_name_whitespace(value)
            else:
                mapped[field] = (value or "").strip() or None
        except ValueError as exc:
            mapped[field] = None
            errors.append(str(exc))

    if config:
        for master_type, field in (
            ("brand", "brand"), ("category", "category"), ("unit", "unit_name"),
        ):
            value = mapped.get(field)
            if value and value not in set(config.get(master_type, [])):
                warnings.append(f"{field}“{value}”未在当前文件配置Master中")
        status_text = (canonical.get("状态") or "").strip()
        if status_text and status_text not in set(config.get("product_status", [])):
            warnings.append(f"状态“{status_text}”未在当前文件配置Master中")
        points_text = (canonical.get("启用积分") or "").strip()
        if points_text and points_text not in set(config.get("points_status", [])):
            warnings.append(f"积分状态“{points_text}”未在当前文件配置Master中")
    if formula_cache_missing:
        warnings.append("公式缺少缓存值：" + "、".join(formula_cache_missing))
    return mapped, warnings, errors


def _json_default(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    raise TypeError(f"{type(value).__name__} cannot be serialized")


def _json_dumps(value) -> str:
    return json.dumps(value, ensure_ascii=False, default=_json_default)


def _existing_target(session: Session, mapped: dict) -> tuple[Product | None, list[dict]]:
    barcode = mapped.get("jan")
    code = mapped.get("qinsi_product_code")
    by_code = session.scalar(select(Product).where(Product.qinsi_product_code == code)) if code else None
    conflicts: list[dict] = []
    barcode_resolution = resolve_local_product_by_jan(session, barcode) if barcode else None
    if barcode_resolution and barcode_resolution.is_conflict:
        conflicts.append({
            "field": "条码",
            "excel_value": barcode,
            "existing_value": [
                product_image_summary(product)
                for product in barcode_resolution.candidate_products
            ],
            "message": "条码在本地对应多个商品，禁止自动选择或新增",
        })
    by_barcode = barcode_resolution.product if barcode_resolution else None

    derived_jan = derive_jan_from_qinsi_sku(code)
    derived_resolution = (
        resolve_local_product_by_jan(session, derived_jan)
        if derived_jan else None
    )
    if derived_resolution and derived_resolution.is_conflict:
        conflicts.append({
            "field": "货号派生条码",
            "excel_value": code,
            "existing_value": [
                product_image_summary(product)
                for product in derived_resolution.candidate_products
            ],
            "message": "货号派生条码在本地对应多个商品，需人工处理",
        })
    by_derived = derived_resolution.product if derived_resolution else None
    matched_products = {
        product.id: product
        for product in (by_code, by_barcode, by_derived)
        if product is not None
    }
    if len(matched_products) > 1:
        conflicts.append({
            "field": "条码/货号",
            "excel_value": {"货号": code, "条码": barcode},
            "existing_value": {
                "命中商品": [product_image_summary(product) for product in matched_products.values()],
            },
            "message": "条码、货号或货号派生条码分别指向不同商品",
        })
        return None, conflicts
    if conflicts:
        return None, conflicts
    target = by_code or by_barcode or by_derived
    if target and barcode and target.jan and target.jan != barcode:
        conflicts.append({
            "field": "条码", "excel_value": barcode, "existing_value": target.jan,
            "product_id": target.id, "message": "货号命中商品，但条码与现有系统值冲突",
        })
    if target and code and target.qinsi_product_code and target.qinsi_product_code != code:
        conflicts.append({
            "field": "货号", "excel_value": code, "existing_value": target.qinsi_product_code,
            "product_id": target.id, "message": "条码命中商品，但货号与现有系统值冲突",
        })
    return (None if conflicts else target), conflicts


def _values_differ(current, incoming) -> bool:
    if isinstance(incoming, Decimal):
        return current is None or Decimal(str(current)) != incoming
    return current != incoming


def _changed_product_fields(target: Product, mapped: dict) -> list[str]:
    changed = [
        field for field in PRODUCT_FIELDS
        if mapped.get(field) is not None and _values_differ(getattr(target, field), mapped[field])
    ]
    if mapped.get("image_url") and _values_differ(target.main_image_source_url, mapped["image_url"]):
        changed.append("main_image_source_url")
    if mapped.get("image_url") and (
        target.main_image_path
        or target.local_image_path
        or _values_differ(target.display_image_url, mapped["image_url"])
    ):
        changed.append("display_image_url")
    return changed


def _apply_qinsi_image_authority(product: Product, image_url: str) -> None:
    product.image_url = image_url
    product.main_image_source_url = image_url
    product.qinsi_image_url = image_url
    product.display_image_url = image_url
    product.main_image_path = None
    product.main_image_hash = None
    product.main_image_locked = False
    product.main_image_source_platform = "qinsi"
    product.main_image_downloaded_at = None
    product.local_image_path = None
    product.image_sha256 = None
    product.image_width = None
    product.image_height = None
    product.image_quality = None
    product.image_localization_status = None
    product.image_localization_source_url = None
    product.image_localized_at = None
    product.image_localization_error = None


def _reset_batch_counts(batch: QinsiImportBatch) -> None:
    for field in (
        "total_rows", "new_count", "update_count", "unchanged_count", "skipped_count",
        "conflict_count", "error_count", "warning_count", "success_count",
    ):
        setattr(batch, field, 0)


def create_import_batch(
    session: Session,
    filename: str,
    content: bytes,
    *,
    business_batch_key: str | None = None,
) -> tuple[QinsiImportBatch, bool]:
    _validate_archive(content)
    file_hash = hashlib.sha256(content).hexdigest()
    existing = session.scalar(select(QinsiImportBatch).where(QinsiImportBatch.file_hash == file_hash))
    if existing is not None:
        return existing, True
    batch = QinsiImportBatch(
        business_batch_key=(business_batch_key or "").strip() or None,
        original_filename=Path(filename).name[:255],
        file_hash=file_hash,
        file_content=content,
        status="queued",
        parse_version=PARSE_VERSION,
    )
    session.add(batch)
    session.commit()
    session.refresh(batch)
    return batch, False


def parse_import_batch(session: Session, batch: QinsiImportBatch) -> QinsiImportBatch:
    if batch.status not in {"queued", "failed", "parsing"}:
        return batch
    batch.status = "parsing"
    batch.error_message = None
    _reset_batch_counts(batch)
    session.execute(delete(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == batch.id))
    session.flush()
    try:
        workbook = parse_qinsi_workbook(batch.file_content)
        prepared: list[tuple[QinsiGoodsImportRow, dict]] = []
        barcode_rows: dict[str, list[QinsiGoodsImportRow]] = {}
        code_rows: dict[str, list[QinsiGoodsImportRow]] = {}
        effective_jan_rows: dict[str, list[QinsiGoodsImportRow]] = {}
        for row_number, raw in workbook.rows:
            mapped, warnings, errors = _mapped_row(
                raw, workbook.config, workbook.missing_formula_cache.get(row_number),
            )
            conflicts: list[dict] = []
            if not any(mapped.get(key) for key in ("name_cn", "qinsi_product_code", "jan")):
                validation_status = "skipped"
                errors = ["空白模板行"]
            elif not mapped.get("name_cn"):
                validation_status = "error"
                errors.append("名称（必填）为空")
            elif not mapped.get("qinsi_product_code"):
                validation_status = "error"
                errors.append("货号（必填且唯一）为空")
            elif errors:
                validation_status = "error"
            else:
                target, conflicts = _existing_target(session, mapped)
                if conflicts:
                    validation_status = "conflict"
                elif target is None:
                    validation_status = "new"
                else:
                    validation_status = (
                        "update" if _changed_product_fields(target, mapped) else "unchanged"
                    )
            row = QinsiGoodsImportRow(
                import_batch_id=batch.id,
                source_file_name=batch.original_filename,
                sheet_name=workbook.sheet_name,
                excel_row_number=row_number,
                qinsi_product_code=mapped.get("qinsi_product_code"),
                barcode=mapped.get("jan"),
                parsed_data=_json_dumps(mapped),
                raw_json=_json_dumps(raw),
                validation_status=validation_status,
                warnings=_json_dumps(warnings) if warnings else None,
                errors=_json_dumps(errors) if errors else None,
                conflict_json=_json_dumps(conflicts) if conflicts else None,
            )
            session.add(row)
            prepared.append((row, mapped))
            if mapped.get("jan"):
                barcode_rows.setdefault(mapped["jan"], []).append(row)
            if mapped.get("qinsi_product_code"):
                code_rows.setdefault(mapped["qinsi_product_code"], []).append(row)
            effective_jans = {
                candidate for candidate in (
                    mapped.get("jan"),
                    mapped.get("qinsi_product_code") if is_valid_jan(mapped.get("qinsi_product_code")) else None,
                    derive_jan_from_qinsi_sku(mapped.get("qinsi_product_code")),
                )
                if candidate
            }
            for effective_jan in effective_jans:
                effective_jan_rows.setdefault(effective_jan, []).append(row)
        session.flush()

        for label, groups in (("条码", barcode_rows), ("货号", code_rows)):
            for value, duplicates in groups.items():
                if len(duplicates) <= 1:
                    continue
                row_numbers = [row.excel_row_number for row in duplicates]
                for row in duplicates:
                    row.validation_status = "conflict"
                    row.conflict_json = _json_dumps([{
                        "field": label,
                        "excel_row_number": row.excel_row_number,
                        "excel_value": value,
                        "existing_value": f"Excel第{','.join(map(str, row_numbers))}行",
                        "message": f"Excel内重复非空{label}：{value}",
                    }])

        for jan, duplicates in effective_jan_rows.items():
            if len(duplicates) <= 1:
                continue
            row_numbers = [row.excel_row_number for row in duplicates]
            for row in duplicates:
                existing_conflicts = json.loads(row.conflict_json or "[]")
                existing_conflicts.append({
                    "field": "统一JAN",
                    "excel_row_number": row.excel_row_number,
                    "excel_value": jan,
                    "existing_value": f"Excel第{','.join(map(str, row_numbers))}行",
                    "message": f"Excel内条码/秦丝货号/合法派生值共同命中 JAN：{jan}",
                })
                row.validation_status = "conflict"
                row.conflict_json = _json_dumps(existing_conflicts)

        statuses = [row.validation_status for row, _ in prepared]
        batch.total_rows = len(prepared)
        batch.new_count = statuses.count("new")
        batch.update_count = statuses.count("update")
        batch.unchanged_count = statuses.count("unchanged")
        batch.skipped_count = statuses.count("skipped")
        batch.conflict_count = statuses.count("conflict")
        batch.error_count = statuses.count("error")
        batch.warning_count = sum(bool(row.warnings) for row, _ in prepared)
        batch.summary_json = _json_dumps({
            "sheet": workbook.sheet_name,
            "source_format": workbook.source_format,
            "headers": list(workbook.headers),
            "configuration": workbook.config,
            "configuration_counts": {key: len(value) for key, value in workbook.config.items()},
            "configuration_method": "配置Sheet对应行实际非空单元格",
            "defined_names": list(workbook.defined_names),
            "inventory_warehouse": workbook.inventory_warehouse,
            "file_sha256": batch.file_hash,
            "parse_version": PARSE_VERSION,
        })
        batch.status = "previewed"
        batch.parse_version = PARSE_VERSION
        session.commit()
        session.refresh(batch)
        return batch
    except Exception as exc:
        session.rollback()
        failed = session.get(QinsiImportBatch, batch.id)
        if failed is not None:
            failed.status = "failed"
            failed.error_message = str(exc)
            session.commit()
        if isinstance(exc, ValueError):
            raise
        raise ValueError(f"秦丝Excel解析失败：{exc}") from exc


def create_import_preview(
    session: Session,
    filename: str,
    content: bytes,
    *,
    business_batch_key: str | None = None,
) -> QinsiImportBatch:
    batch, reused = create_import_batch(
        session, filename, content, business_batch_key=business_batch_key,
    )
    if reused and batch.status not in {"failed", "queued", "parsing"}:
        return batch
    return parse_import_batch(session, batch)


def process_queued_import(database_url: str, batch_id: int) -> None:
    engine = build_engine(database_url)
    try:
        with Session(engine) as session:
            batch = session.get(QinsiImportBatch, batch_id)
            if batch is not None:
                parse_import_batch(session, batch)
    finally:
        engine.dispose()


def _normalize_master_name(value: str) -> str:
    return unicodedata.normalize("NFKC", value).casefold().strip()


def _sync_masters(
    session: Session,
    batch: QinsiImportBatch,
    configuration: dict[str, list[str]],
    now: datetime,
) -> dict[str, dict[str, QinsiMasterValue]]:
    result: dict[str, dict[str, QinsiMasterValue]] = {}
    for master_type, values in configuration.items():
        typed: dict[str, QinsiMasterValue] = {}
        for source_name in values:
            master = session.scalar(select(QinsiMasterValue).where(
                QinsiMasterValue.source_system == "qinsi",
                QinsiMasterValue.master_type == master_type,
                QinsiMasterValue.source_name == source_name,
            ))
            if master is None:
                master = QinsiMasterValue(
                    master_type=master_type,
                    source_name=source_name,
                    normalized_name=_normalize_master_name(source_name),
                    source_system="qinsi",
                    first_import_batch_id=batch.id,
                    last_seen_at=now,
                )
                session.add(master)
                session.flush()
            else:
                master.is_active = True
                master.last_seen_at = now
            typed[source_name] = master
        result[master_type] = typed
    return result


def _warehouse_code(name: str) -> str:
    return f"QW-QINSI-{hashlib.sha256(name.encode('utf-8')).hexdigest()[:12].upper()}"


def _sync_warehouses(session: Session, names: list[str]) -> dict[str, Location]:
    result: dict[str, Location] = {}
    locations = list(session.scalars(select(Location)))
    next_sort = max((location.sort_order for location in locations), default=0)
    for name in names:
        location = next((item for item in locations if item.display_name == name), None)
        if location is not None and not location.is_qinsi_warehouse:
            raise ValueError(f"配置仓库“{name}”与现有非秦丝位置同名，请先人工处理")
        if location is None:
            next_sort += 10
            location = Location(
                internal_code=_warehouse_code(name),
                display_name=name,
                location_type="qinsi_warehouse",
                is_qinsi_warehouse=True,
                is_active=True,
                sort_order=next_sort,
            )
            session.add(location)
            session.flush()
            locations.append(location)
        else:
            location.is_active = True
        result[name] = location
    return result


def _decoded_mapping(row: QinsiGoodsImportRow) -> dict:
    mapped = json.loads(row.parsed_data)
    for field in (
        "purchase_price", "sale_price", "minimum_sale_price",
        "inventory_warning_lower", "inventory_warning_upper", "weight_kg",
    ):
        if mapped.get(field) is not None:
            mapped[field] = Decimal(str(mapped[field]))
    return mapped


def _ensure_product_barcode(session: Session, product: Product, barcode: str | None) -> None:
    if not barcode:
        return
    existing = session.scalar(select(ProductBarcode).where(ProductBarcode.barcode == barcode))
    if existing is not None and existing.product_id != product.id:
        raise ValueError(f"条码 {barcode} 已关联其他商品")
    if existing is None:
        session.add(ProductBarcode(
            product_id=product.id, barcode=barcode, source_system="qinsi", is_primary=True,
        ))


def _build_inventory_snapshot(
    session: Session,
    batch: QinsiImportBatch,
    rows: list[QinsiGoodsImportRow],
    warehouse_name: str | None,
    warehouses: dict[str, Location],
    now: datetime,
) -> QinsiInventorySnapshot | None:
    inventory_rows: list[tuple[QinsiGoodsImportRow, dict]] = []
    for row in rows:
        if row.validation_status in {"conflict", "error", "skipped"}:
            continue
        mapped = _decoded_mapping(row)
        if (
            mapped.get("counted_inventory_quantity") is not None
            or mapped.get("current_inventory_quantity") is not None
        ):
            inventory_rows.append((row, mapped))
    if not inventory_rows:
        return None
    existing = session.scalar(select(QinsiInventorySnapshot).where(
        (QinsiInventorySnapshot.source_import_batch_id == batch.id)
        | (QinsiInventorySnapshot.file_hash == batch.file_hash)
    ))
    if existing is not None:
        if existing.source_import_batch_id is None:
            existing.source_import_batch_id = batch.id
        return existing

    warehouse = warehouses.get(warehouse_name or "")
    snapshot = QinsiInventorySnapshot(
        batch_no=f"QSG-{now:%Y%m%d}-{uuid.uuid4().hex[:10].upper()}",
        original_filename=batch.original_filename,
        file_hash=batch.file_hash,
        file_content=batch.file_content,
        imported_at=now,
        data_at=now,
        source_system="qinsi",
        snapshot_type="counted_inventory",
        source_import_batch_id=batch.id,
        status="completed",
    )
    session.add(snapshot)
    session.flush()
    for row, mapped in inventory_rows:
        raw = json.loads(row.raw_json)
        errors: list[str] = []
        if warehouse is None:
            errors.append(f"未知秦丝仓库：{warehouse_name or '未提供'}")
        if row.product_id is None:
            errors.append("未匹配到本地商品")
        current = mapped.get("current_inventory_quantity")
        line = QinsiInventorySnapshotLine(
            snapshot_id=snapshot.id,
            original_row_no=row.excel_row_number,
            raw_product_name=mapped.get("name_cn"),
            jan=mapped.get("jan"),
            qinsi_product_code=mapped.get("qinsi_product_code"),
            raw_warehouse_name=warehouse_name,
            quantity=mapped.get("counted_inventory_quantity"),
            current_quantity=Decimal(current).quantize(Decimal("0.001")) if current is not None else None,
            raw_summary_json=_json_dumps(raw),
            product_id=row.product_id,
            warehouse_id=warehouse.id if warehouse else None,
            matching_method="qinsi_product_code" if row.product_id else None,
            matching_status="matched" if row.product_id else "unmatched",
            warehouse_status="matched" if warehouse else "unknown",
            error_message="；".join(errors) or None,
        )
        snapshot.lines.append(line)
    snapshot.total_rows = len(snapshot.lines)
    snapshot.success_rows = sum(
        line.product_id is not None and line.warehouse_id is not None and line.quantity is not None
        for line in snapshot.lines
    )
    snapshot.unmatched_rows = sum(line.product_id is None for line in snapshot.lines)
    snapshot.exception_rows = sum(
        line.warehouse_id is None or line.quantity is None for line in snapshot.lines
    )
    snapshot.status = (
        "completed_with_issues"
        if snapshot.unmatched_rows or snapshot.exception_rows else "completed"
    )
    snapshot.error_summary = "；".join(dict.fromkeys(
        line.error_message for line in snapshot.lines if line.error_message
    )) or None
    return snapshot


def confirm_import(
    session: Session,
    batch: QinsiImportBatch,
    *,
    allow_importing: bool = False,
) -> QinsiImportBatch:
    allowed = {"previewed", "importing"} if allow_importing else {"previewed"}
    if batch.status not in allowed:
        raise ValueError("导入任务状态不允许确认；修复前任务必须重新解析或重新上传")
    rows = list(session.scalars(select(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
    ).order_by(QinsiGoodsImportRow.excel_row_number)))
    summary = json.loads(batch.summary_json or "{}")
    configuration = summary.get("configuration") or {}
    now = datetime.now(timezone.utc)
    imported_new = imported_update = unchanged = 0
    try:
        masters = _sync_masters(session, batch, configuration, now)
        warehouses = _sync_warehouses(session, configuration.get("warehouse", []))
        for row in rows:
            if row.validation_status in {"skipped", "conflict", "error"}:
                continue
            mapped = _decoded_mapping(row)
            target, conflicts = _existing_target(session, mapped)
            if conflicts:
                row.validation_status = "conflict"
                row.conflict_json = _json_dumps(conflicts)
                continue
            if target is None:
                target = Product(product_origin="qinsi", source="qinsi_import")
                session.add(target)
                action = "imported_new"
            else:
                action = "imported_updated" if _changed_product_fields(target, mapped) else "unchanged"
            for field in PRODUCT_FIELDS:
                value = mapped.get(field)
                if value is not None and _values_differ(getattr(target, field), value):
                    setattr(target, field, value)
            if mapped.get("name_cn"):
                target.name_source = "qinsi_import"
                target.name_locked = False
                target.product_data_confirmed = False
                if mapped.get("name_ja") is None:
                    target.name_ja = None
            if mapped.get("image_url"):
                _apply_qinsi_image_authority(target, mapped["image_url"])
            target.display_name = format_product_display_name(target.name_cn, target.name_ja)
            target.status = "qinsi_product_imported"
            for master_type, (field, relationship_field) in MASTER_PRODUCT_FIELDS.items():
                value = mapped.get(field)
                master = masters.get(master_type, {}).get(value)
                if master is not None:
                    setattr(target, relationship_field, master.id)
            session.flush()
            if mapped.get("jan"):
                merge_products_with_jan_into_primary(
                    session,
                    mapped["jan"],
                    target,
                    actor="qinsi_import",
                    commit=False,
                )
                merge_duplicate_jan_group(
                    session,
                    mapped["jan"],
                    primary_product_id=target.id,
                    actor="qinsi_import",
                    commit=False,
                )
            _ensure_product_barcode(session, target, mapped.get("jan"))
            ensure_qinsi_derived_barcode(session, target)
            if mapped.get("image_url"):
                queue_product_image_localization(session, target)
            row.product_id = target.id
            row.validation_status = action
            if action == "imported_new":
                imported_new += 1
            elif action == "imported_updated":
                imported_update += 1
            else:
                unchanged += 1

        _build_inventory_snapshot(
            session, batch, rows, summary.get("inventory_warehouse"), warehouses, now,
        )
        batch.new_count = imported_new
        batch.update_count = imported_update
        batch.unchanged_count = unchanged
        batch.success_count = imported_new + imported_update
        batch.skipped_count = sum(row.validation_status == "skipped" for row in rows)
        batch.conflict_count = sum(row.validation_status == "conflict" for row in rows)
        batch.error_count = sum(row.validation_status == "error" for row in rows)
        batch.warning_count = sum(bool(row.warnings) for row in rows)
        batch.status = (
            "completed_with_issues" if batch.conflict_count or batch.error_count else "completed"
        )
        batch.confirmed_at = now
        session.commit()
    except Exception:
        session.rollback()
        raise
    session.refresh(batch)
    return batch


def queue_import_confirmation(session: Session, batch: QinsiImportBatch) -> None:
    if batch.status != "previewed":
        raise ValueError("导入任务状态不允许确认")
    batch.status = "importing"
    session.commit()


def process_queued_confirmation(database_url: str, batch_id: int) -> None:
    engine = build_engine(database_url)
    try:
        with Session(engine) as session:
            batch = session.get(QinsiImportBatch, batch_id)
            if batch is not None:
                try:
                    confirm_import(session, batch, allow_importing=True)
                except Exception as exc:
                    session.rollback()
                    failed = session.get(QinsiImportBatch, batch_id)
                    if failed is not None:
                        failed.status = "failed"
                        failed.error_message = str(exc)
                        session.commit()
    finally:
        engine.dispose()
