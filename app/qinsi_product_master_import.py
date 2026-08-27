from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.db import build_engine
from app.local_product import is_valid_jan
from app.models import (
    Product,
    ProductBarcode,
    ProductOperationLog,
    QinsiConflictResolution,
    QinsiGoodsImportRow,
    QinsiImportBatch,
    QinsiProductMapping,
)
from app.product_merge import merge_duplicate_jan_group, merge_products_with_jan_into_primary
from app.product_identity import format_product_display_name, normalize_product_name_whitespace


MAX_XLSX_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MASTER_PARSE_VERSION = 4
MASTER_SOURCE = "qinsi_product_master"
MASTER_SHEET = "Sheet1"
QINSI_CONFLICT_RESOLUTION_TYPES = {
    "shared_barcode_variant": {
        "label": "颜色/规格不同，共用同一条码",
        "action": "keep_multiple_products_share_barcode",
    },
    "qinsi_legacy_data": {
        "label": "秦丝历史旧数据",
        "action": "keep_qinsi_entity_sync_current_master",
    },
    "code_barcode_overlap": {
        "label": "货号与单品条码角色重叠",
        "action": "allow_product_code_and_barcode_overlap",
    },
    "true_duplicate": {
        "label": "同一商品重复登记",
        "action": "merge_or_link_to_primary_product",
    },
    "wrong_barcode": {
        "label": "条码录入错误",
        "action": "keep_qinsi_record_save_manual_mapping",
    },
    "other": {
        "label": "其他",
        "action": "allow_with_warning",
    },
}
NON_MERGE_RESOLUTION_TYPES = {
    "shared_barcode_variant",
    "qinsi_legacy_data",
    "code_barcode_overlap",
    "wrong_barcode",
    "other",
}

REQUIRED_HEADERS = (
    "商品名称", "货号", "商品条码", "单品条码", "图片链接", "品牌", "分类", "单位",
    "采购价", "销售价", "状态", "备注",
)
HEADER_TO_FIELD = {
    "商品名称": "qinsi_name",
    "商品规格": "specification",
    "货号": "qinsi_goods_no",
    "商品条码": "qinsi_product_barcode",
    "单品条码": "qinsi_unit_barcode",
    "型号规格": "model_spec",
    "图片链接": "qinsi_image_url",
    "品牌": "qinsi_brand",
    "分类": "qinsi_category",
    "单位": "qinsi_unit",
    "采购价": "qinsi_purchase_price",
    "销售价": "qinsi_sale_price",
    "最低销售价": "minimum_sale_price",
    "保质期": "shelf_life_days",
    "产地": "origin_place",
    "适用年龄": "applicable_age",
    "排序": "qinsi_sort_order",
    "状态": "qinsi_status",
    "库存预警下限": "inventory_warning_lower",
    "库存预警上限": "inventory_warning_upper",
    "备注": "qinsi_remark",
}
MONEY_FIELDS = {"qinsi_purchase_price", "qinsi_sale_price", "minimum_sale_price"}
INTEGER_FIELDS = {"shelf_life_days", "qinsi_sort_order"}
DECIMAL3_FIELDS = {"inventory_warning_lower", "inventory_warning_upper"}
SCIENTIFIC_IDENTIFIER_PATTERN = re.compile(r"^[+-]?\d+(?:\.\d+)?[eE][+-]?\d+$")


@dataclass(frozen=True, slots=True)
class MasterInputFile:
    filename: str
    content: bytes


def _validate_archive(content: bytes) -> None:
    if not content or len(content) > MAX_XLSX_BYTES:
        raise ValueError("Excel文件为空或超过20MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(item.file_size for item in archive.infolist()) > MAX_UNCOMPRESSED_BYTES
            ):
                raise ValueError("Excel解压内容过大")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("文件不是有效的 .xlsx 工作簿")
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 .xlsx 工作簿") from exc


def _plain_decimal(value: int | float | Decimal) -> str:
    try:
        number = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError(f"Excel数值无效：{value}") from exc
    if not number.is_finite():
        raise ValueError(f"Excel数值无效：{value}")
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _zero_format_width(number_format: str) -> int | None:
    first = (number_format or "").split(";", 1)[0]
    return len(first) if first and set(first) == {"0"} else None


def _logical_cell_text(cell: Cell, *, identifier: bool = False) -> str | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        text = _plain_decimal(value)
        if identifier:
            width = _zero_format_width(cell.number_format)
            if width and text.isdigit():
                text = text.zfill(width)
        return text
    text = str(value).strip()
    return text or None


def _identifier(value: str | None, label: str) -> str | None:
    text = (value or "").strip()
    if not text:
        return None
    if SCIENTIFIC_IDENTIFIER_PATTERN.fullmatch(text):
        raise ValueError(f"{label}使用了科学计数法，无法保证原始标识符")
    return text


def _decimal_value(value: str | None, label: str, *, scale: int) -> Decimal | None:
    text = (value or "").strip()
    if not text:
        return None
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"{label}不是有效数值") from exc
    if not number.is_finite():
        raise ValueError(f"{label}不是有效数值")
    quantum = Decimal(1).scaleb(-scale)
    rounded = number.quantize(quantum)
    if rounded != number:
        raise ValueError(f"{label}最多支持{scale}位小数")
    return rounded


def _integer_value(value: str | None, label: str) -> int | None:
    number = _decimal_value(value, label, scale=6)
    if number is None:
        return None
    if number != number.to_integral_value():
        raise ValueError(f"{label}必须是整数")
    return int(number)


def _json_default(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    raise TypeError(f"{type(value).__name__} cannot be serialized")


def _json_dumps(value) -> str:
    return json.dumps(value, ensure_ascii=False, default=_json_default)


def _json_loads(value: str | None, fallback):
    if not value:
        return fallback
    try:
        return json.loads(value)
    except (TypeError, ValueError):
        return fallback


def _combined_hash(files: list[MasterInputFile]) -> str:
    digest = hashlib.sha256()
    for item in files:
        digest.update(item.filename.encode("utf-8"))
        digest.update(b"\0")
        digest.update(hashlib.sha256(item.content).hexdigest().encode("ascii"))
        digest.update(b"\0")
    return digest.hexdigest()


def _manifest_bytes(files: list[MasterInputFile]) -> bytes:
    return _json_dumps({
        "kind": MASTER_SOURCE,
        "files": [
            {"filename": item.filename, "sha256": hashlib.sha256(item.content).hexdigest(), "bytes": len(item.content)}
            for item in files
        ],
    }).encode("utf-8")


def determine_qinsi_master_jan(
    *,
    unit_barcode: str | None,
    product_barcode: str | None,
    goods_no: str | None,
) -> tuple[str | None, str]:
    for source, value in (
        ("unit_barcode", unit_barcode),
        ("product_barcode", product_barcode),
        ("goods_no", goods_no),
    ):
        if is_valid_jan(value):
            return value, source
    return None, "none"


def parse_qinsi_master_files(files: list[MasterInputFile]) -> list[dict]:
    rows: list[dict] = []
    for item in files:
        _validate_archive(item.content)
        try:
            workbook = load_workbook(io.BytesIO(item.content), data_only=True, read_only=False, keep_links=False)
        except Exception as exc:
            raise ValueError(f"{item.filename} 无法解析") from exc
        try:
            sheet = workbook.active
            headers = [_logical_cell_text(cell) or "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            missing = [header for header in REQUIRED_HEADERS if header not in headers]
            if missing:
                raise ValueError(f"{item.filename} 缺少表头：{'、'.join(missing)}")
            for excel_row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
                raw = {headers[index]: _logical_cell_text(cell, identifier=headers[index] in {"货号", "商品条码", "单品条码"})
                       for index, cell in enumerate(cells) if index < len(headers)}
                if not any(raw.values()):
                    continue
                mapped: dict = {
                    "source_file_name": item.filename,
                    "excel_row_number": excel_row_number,
                }
                warnings: list[str] = []
                errors: list[str] = []
                for header, field in HEADER_TO_FIELD.items():
                    value = raw.get(header)
                    try:
                        if field in {"qinsi_goods_no", "qinsi_product_barcode", "qinsi_unit_barcode"}:
                            mapped[field] = _identifier(value, header)
                        elif field in MONEY_FIELDS:
                            mapped[field] = _decimal_value(value, header, scale=2)
                        elif field in DECIMAL3_FIELDS:
                            mapped[field] = _decimal_value(value, header, scale=3)
                        elif field in INTEGER_FIELDS:
                            mapped[field] = _integer_value(value, header)
                        else:
                            mapped[field] = normalize_product_name_whitespace(value) if field == "qinsi_name" else ((value or "").strip() or None)
                    except ValueError as exc:
                        mapped[field] = None
                        errors.append(str(exc))
                jan, jan_source = determine_qinsi_master_jan(
                    unit_barcode=mapped.get("qinsi_unit_barcode"),
                    product_barcode=mapped.get("qinsi_product_barcode"),
                    goods_no=mapped.get("qinsi_goods_no"),
                )
                mapped["jan"] = jan
                mapped["jan_source"] = jan_source
                mapped["has_jan"] = jan is not None
                mapped["warnings"] = warnings
                mapped["errors"] = errors
                mapped["raw"] = raw
                rows.append(mapped)
        finally:
            workbook.close()
    return rows


def _products_by_jan(session: Session, jan: str) -> list[Product]:
    return list(session.scalars(select(Product).where(Product.jan == jan).order_by(Product.id)))


def _products_by_qinsi_goods_no(session: Session, goods_no: str) -> list[Product]:
    product_ids: set[int] = set()
    product_ids.update(
        session.scalars(select(Product.id).where(Product.qinsi_product_code == goods_no)).all()
    )
    product_ids.update(
        session.scalars(select(QinsiProductMapping.product_id).where(QinsiProductMapping.qinsi_product_code == goods_no)).all()
    )
    if not product_ids:
        return []
    return list(session.scalars(select(Product).where(Product.id.in_(product_ids)).order_by(Product.id)))


def _qinsi_codes_for_barcode_group(rows: list[dict], barcode: str | None) -> list[str]:
    if not barcode:
        return []
    return sorted({
        str(item.get("qinsi_goods_no"))
        for item in rows
        if item.get("jan") == barcode and item.get("qinsi_goods_no")
    })


def qinsi_conflict_key(
    *,
    barcode: str | None,
    qinsi_product_codes: list[str] | tuple[str, ...],
    conflict_type: str,
) -> str:
    code_part = ",".join(sorted({str(code) for code in qinsi_product_codes if code}))
    return f"{barcode or ''}|{code_part}|{conflict_type}"


def _resolution_key_for_row(mapped: dict, rows: list[dict], resolution_type: str) -> str:
    return qinsi_conflict_key(
        barcode=mapped.get("jan"),
        qinsi_product_codes=_qinsi_codes_for_barcode_group(rows, mapped.get("jan")) or [mapped.get("qinsi_goods_no")],
        conflict_type=resolution_type,
    )


def _auto_resolution_for_row(
    session: Session,
    mapped: dict,
    rows: list[dict],
) -> QinsiConflictResolution | None:
    keys = [
        _resolution_key_for_row(mapped, rows, resolution_type)
        for resolution_type in QINSI_CONFLICT_RESOLUTION_TYPES
    ]
    return session.scalar(
        select(QinsiConflictResolution)
        .where(
            QinsiConflictResolution.conflict_key.in_(keys),
            QinsiConflictResolution.auto_apply.is_(True),
        )
        .order_by(QinsiConflictResolution.updated_at.desc(), QinsiConflictResolution.id.desc())
        .limit(1)
    )


def _annotate_resolution(mapped: dict, resolution: QinsiConflictResolution | None, *, manual: bool = False) -> None:
    if resolution is None:
        return
    mapped["qinsi_conflict_resolution_type"] = resolution.resolution_type
    mapped["qinsi_conflict_resolution_key"] = resolution.conflict_key
    mapped["qinsi_conflict_resolution_action"] = resolution.action
    mapped["qinsi_conflict_resolution_note"] = resolution.note
    mapped["qinsi_conflict_resolution_auto_applied"] = not manual
    if resolution.resolution_type == "shared_barcode_variant":
        mapped["qinsi_shared_barcode"] = mapped.get("jan")
        mapped["force_no_product_jan"] = True
    warnings = list(mapped.get("warnings") or [])
    label = QINSI_CONFLICT_RESOLUTION_TYPES.get(resolution.resolution_type, {}).get("label", resolution.resolution_type)
    prefix = "已按历史规则处理" if not manual else "已按人工确认处理"
    warning = f"{prefix}：{label}"
    if warning not in warnings:
        warnings.append(warning)
    mapped["warnings"] = warnings


def _resolution_allows_row(mapped: dict) -> bool:
    return mapped.get("qinsi_conflict_resolution_type") in NON_MERGE_RESOLUTION_TYPES


def _product_summary(product: Product) -> dict:
    return {
        "product_id": product.id,
        "internal_sku": product.internal_sku,
        "jan": product.jan,
        "qinsi_goods_no": product.qinsi_product_code,
        "name": product.display_name or product.name_cn or product.name_ja,
    }


def _changed_product_fields(target: Product, mapped: dict) -> list[str]:
    fields = [
        "qinsi_product_code", "qinsi_product_barcode", "qinsi_unit_barcode",
        "qinsi_name", "qinsi_image_url", "qinsi_brand", "qinsi_category",
        "qinsi_unit", "qinsi_status", "qinsi_remark", "purchase_price",
        "sale_price", "minimum_sale_price", "brand", "category", "unit_name",
        "model_spec", "specification", "origin_place", "applicable_age",
        "qinsi_sort_order", "inventory_warning_lower", "inventory_warning_upper",
        "shelf_life_days", "has_jan", "source", "product_origin", "status",
    ]
    changed: list[str] = []
    if target.jan != mapped.get("jan"):
        changed.append("jan")
    for field in fields:
        incoming = _incoming_value_for_field(field, mapped)
        if incoming is not None and getattr(target, field, None) != incoming:
            changed.append(field)
    if mapped.get("qinsi_name") and target.name_cn != mapped.get("qinsi_name"):
        changed.append("name_cn")
    if mapped.get("qinsi_image_url") and (
        target.image_url != mapped.get("qinsi_image_url")
        or target.main_image_source_url != mapped.get("qinsi_image_url")
        or target.display_image_url != mapped.get("qinsi_image_url")
        or target.main_image_path
        or target.local_image_path
    ):
        changed.append("display_image_url")
    return changed


def _incoming_value_for_field(field: str, mapped: dict):
    mapping = {
        "qinsi_product_code": "qinsi_goods_no",
        "qinsi_product_barcode": "qinsi_product_barcode",
        "qinsi_unit_barcode": "qinsi_unit_barcode",
        "qinsi_name": "qinsi_name",
        "qinsi_image_url": "qinsi_image_url",
        "qinsi_brand": "qinsi_brand",
        "qinsi_category": "qinsi_category",
        "qinsi_unit": "qinsi_unit",
        "qinsi_status": "qinsi_status",
        "qinsi_remark": "qinsi_remark",
        "purchase_price": "qinsi_purchase_price",
        "sale_price": "qinsi_sale_price",
        "brand": "qinsi_brand",
        "category": "qinsi_category",
        "unit_name": "qinsi_unit",
        "has_jan": "has_jan",
    }
    if field == "source":
        return "qinsi_import"
    if field == "product_origin":
        return "qinsi"
    if field == "status":
        return "qinsi_product_imported"
    return mapped.get(mapping.get(field, field))


def _looks_like_placeholder(value: str | None) -> bool:
    text = (value or "").strip()
    return not text or any(token in text for token in ("缺商品", "待补", "未命名", "中文名待补", "日文名待补"))


def _should_fill_name(product: Product, mapped: dict) -> bool:
    if not mapped.get("qinsi_name"):
        return False
    if product.name_locked or product.product_data_confirmed:
        return False
    if product.name_source in {"qinsi_import", "receipt", "auto", "missing", None} and (
        _looks_like_placeholder(product.name_cn)
        or _looks_like_placeholder(product.display_name)
        or product.status in {"new_pending_completion", "new_pending_review", "pending_qinsi_product_import"}
    ):
        return True
    return product.name_cn is None and product.name_ja is None


def _apply_qinsi_image_authority(product: Product, image_url: str) -> None:
    product.image_url = image_url
    product.main_image_source_url = image_url
    product.qinsi_image_url = image_url
    product.display_image_url = image_url
    product.main_image_path = None
    product.main_image_hash = None
    product.main_image_locked = False
    product.main_image_source_platform = "qinsi"
    product.main_image_downloaded_at = None
    product.local_image_path = None
    product.image_sha256 = None
    product.image_width = None
    product.image_height = None
    product.image_quality = None
    product.image_localization_status = None
    product.image_localization_source_url = None
    product.image_localized_at = None
    product.image_localization_error = None


def _apply_product_fields(session: Session, product: Product, mapped: dict, now: datetime) -> None:
    old_jan = product.jan
    new_jan = None if mapped.get("force_no_product_jan") else mapped.get("jan")
    if old_jan != new_jan:
        session.add(ProductOperationLog(
            product_id=product.id,
            internal_sku=product.internal_sku,
            action="edit",
            actor="qinsi_master_import",
            reason="秦丝全量商品主档同步修正JAN",
            before_json=_json_dumps({"jan": old_jan}),
            after_json=_json_dumps({
                "jan": new_jan,
                "qinsi_goods_no": mapped.get("qinsi_goods_no"),
                "source_file_name": mapped.get("source_file_name"),
                "excel_row_number": mapped.get("excel_row_number"),
            }),
        ))
    product.jan = new_jan
    product.qinsi_product_code = mapped.get("qinsi_goods_no")
    product.qinsi_product_barcode = mapped.get("qinsi_product_barcode")
    product.qinsi_unit_barcode = mapped.get("qinsi_unit_barcode")
    product.qinsi_name = mapped.get("qinsi_name")
    product.qinsi_image_url = mapped.get("qinsi_image_url")
    product.qinsi_brand = mapped.get("qinsi_brand")
    product.qinsi_category = mapped.get("qinsi_category")
    product.qinsi_unit = mapped.get("qinsi_unit")
    product.qinsi_status = mapped.get("qinsi_status")
    product.qinsi_remark = mapped.get("qinsi_remark")
    product.qinsi_synced_at = now
    product.has_jan = mapped.get("has_jan", False)
    product.source = "qinsi_import"
    product.product_origin = "qinsi"
    product.status = "qinsi_product_imported"
    for product_field, mapped_field in (
        ("purchase_price", "qinsi_purchase_price"),
        ("sale_price", "qinsi_sale_price"),
        ("minimum_sale_price", "minimum_sale_price"),
        ("brand", "qinsi_brand"),
        ("category", "qinsi_category"),
        ("unit_name", "qinsi_unit"),
        ("model_spec", "model_spec"),
        ("specification", "specification"),
        ("origin_place", "origin_place"),
        ("applicable_age", "applicable_age"),
        ("qinsi_sort_order", "qinsi_sort_order"),
        ("inventory_warning_lower", "inventory_warning_lower"),
        ("inventory_warning_upper", "inventory_warning_upper"),
        ("shelf_life_days", "shelf_life_days"),
    ):
        value = mapped.get(mapped_field)
        if value is not None:
            setattr(product, product_field, value)
    if mapped.get("qinsi_image_url"):
        _apply_qinsi_image_authority(product, mapped["qinsi_image_url"])
    if mapped.get("qinsi_name"):
        product.name_cn = mapped["qinsi_name"]
        product.name_ja = None
        product.name_source = "qinsi_import"
        product.name_locked = False
        product.product_data_confirmed = False
        product.display_name = format_product_display_name(product.name_cn, product.name_ja)
    elif not product.display_name:
        product.display_name = format_product_display_name(product.name_cn, product.name_ja)
    session.flush()
    _sync_mapping(session, product, mapped.get("qinsi_goods_no"))
    _sync_product_barcodes(session, product, mapped)


def _sync_product_barcodes(session: Session, product: Product, mapped: dict) -> None:
    for barcode, source in (
        (mapped.get("qinsi_product_barcode"), "qinsi_product_barcode"),
        (mapped.get("qinsi_unit_barcode"), "qinsi_unit_barcode"),
        (mapped.get("qinsi_shared_barcode"), "qinsi_shared_barcode"),
    ):
        if not barcode or not is_valid_jan(barcode):
            continue
        exists = session.scalar(select(ProductBarcode).where(
            ProductBarcode.product_id == product.id,
            ProductBarcode.barcode == barcode,
        ))
        if exists is None:
            session.add(ProductBarcode(
                product_id=product.id,
                barcode=barcode,
                source_system=source,
                is_primary=source != "qinsi_shared_barcode",
            ))


def _sync_mapping(session: Session, product: Product, goods_no: str | None) -> None:
    if not goods_no:
        return
    mapping = session.scalar(select(QinsiProductMapping).where(QinsiProductMapping.qinsi_product_code == goods_no))
    if mapping is None:
        session.add(QinsiProductMapping(qinsi_product_code=goods_no, product_id=product.id))
    else:
        mapping.product_id = product.id


def _prepare_rows(session: Session, rows: list[dict]) -> tuple[list[tuple[dict, str, list[dict], Product | None]], dict]:
    jan_groups: dict[str, list[dict]] = {}
    goods_groups: dict[str, list[dict]] = {}
    for mapped in rows:
        if mapped.get("jan"):
            jan_groups.setdefault(mapped["jan"], []).append(mapped)
        if mapped.get("qinsi_goods_no"):
            goods_groups.setdefault(mapped["qinsi_goods_no"], []).append(mapped)

    prepared: list[tuple[dict, str, list[dict], Product | None]] = []
    counts = {
        "jan_from_unit_barcode": 0,
        "jan_from_product_barcode": 0,
        "jan_from_goods_no": 0,
        "no_jan_count": 0,
        "same_jan_multi_group_count": sum(1 for group in jan_groups.values() if len(group) > 1),
        "same_jan_multi_product_count": sum(len(group) for group in jan_groups.values() if len(group) > 1),
        "local_exists_count": 0,
        "jan_correction_count": 0,
        "mark_imported_count": 0,
    }
    for mapped in rows:
        if mapped["jan_source"] == "unit_barcode":
            counts["jan_from_unit_barcode"] += 1
        elif mapped["jan_source"] == "product_barcode":
            counts["jan_from_product_barcode"] += 1
        elif mapped["jan_source"] == "goods_no":
            counts["jan_from_goods_no"] += 1
        else:
            counts["no_jan_count"] += 1

        resolution = _auto_resolution_for_row(session, mapped, rows)
        _annotate_resolution(mapped, resolution)
        allows_resolved = _resolution_allows_row(mapped)
        conflicts: list[dict] = []
        errors = list(mapped.get("errors") or [])
        if not mapped.get("qinsi_goods_no"):
            errors.append("货号为空，无法建立秦丝身份")
        if not mapped.get("qinsi_name"):
            errors.append("商品名称为空")
        if errors:
            mapped["errors"] = errors
            prepared.append((mapped, "error", conflicts, None))
            continue

        if not allows_resolved and mapped.get("qinsi_goods_no") and len(goods_groups[mapped["qinsi_goods_no"]]) > 1:
            conflicts.append({
                "field": "货号",
                "excel_value": mapped["qinsi_goods_no"],
                "message": f"Excel内重复秦丝货号：{mapped['qinsi_goods_no']}",
            })
        if not allows_resolved and mapped.get("jan") and len(jan_groups[mapped["jan"]]) > 1:
            conflicts.append({
                "field": "判定JAN",
                "excel_value": mapped["jan"],
                "message": f"Excel内同一JAN出现多个秦丝商品：{mapped['jan']}，需先在秦丝或本地合并后再导入",
            })

        target: Product | None = None
        goods_products = _products_by_qinsi_goods_no(session, mapped["qinsi_goods_no"])
        if len(goods_products) > 1:
            conflicts.append({
                "field": "货号",
                "excel_value": mapped["qinsi_goods_no"],
                "existing_value": [_product_summary(product) for product in goods_products],
                "message": "同一qinsi_goods_no对应多个本地Product",
            })
        elif goods_products:
            target = goods_products[0]

        if target is None and mapped.get("jan") and not mapped.get("force_no_product_jan"):
            jan_products = _products_by_jan(session, mapped["jan"])
            unclaimed = [product for product in jan_products if not product.qinsi_product_code]
            if len(jan_products) == 1:
                candidate = jan_products[0]
                if not candidate.qinsi_product_code and (len(jan_groups[mapped["jan"]]) == 1 or allows_resolved):
                    target = candidate
                elif candidate.qinsi_product_code == mapped.get("qinsi_goods_no"):
                    target = candidate
            elif len(jan_products) > 1 and unclaimed:
                conflicts.append({
                    "field": "判定JAN",
                    "excel_value": mapped["jan"],
                    "existing_value": [_product_summary(product) for product in jan_products],
                    "message": "同一JAN存在未绑定秦丝货号的本地Product，无法自动判断对应秦丝商品",
                })

        if conflicts:
            prepared.append((mapped, "conflict", conflicts, None))
            continue
        if target is None:
            prepared.append((mapped, "new", conflicts, None))
        else:
            counts["local_exists_count"] += 1
            if target.jan != (None if mapped.get("force_no_product_jan") else mapped.get("jan")):
                counts["jan_correction_count"] += 1
            status = "update" if _changed_product_fields(target, mapped) else "unchanged"
            prepared.append((mapped, status, conflicts, target))

    counts["mark_imported_count"] = sum(1 for _, status, _, _ in prepared if status in {"new", "update", "unchanged"})
    return prepared, counts


def create_qinsi_master_preview(
    session: Session,
    files: list[MasterInputFile],
    *,
    business_batch_key: str | None = None,
) -> QinsiImportBatch:
    if not files:
        raise ValueError("请至少上传一个Excel文件")
    file_hash = _combined_hash(files)
    existing = session.scalar(select(QinsiImportBatch).where(QinsiImportBatch.file_hash == file_hash))
    if existing is not None and existing.source_system == MASTER_SOURCE:
        if existing.status in {"completed", "completed_with_issues", "importing"}:
            return existing
        session.execute(delete(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == existing.id))
        session.delete(existing)
        session.flush()
    batch = QinsiImportBatch(
        business_batch_key=(business_batch_key or "").strip() or None,
        source_system=MASTER_SOURCE,
        original_filename="; ".join(item.filename for item in files),
        file_hash=file_hash,
        file_content=_manifest_bytes(files),
        status="parsing",
        parse_version=MASTER_PARSE_VERSION,
    )
    session.add(batch)
    session.flush()
    try:
        rows = parse_qinsi_master_files(files)
        prepared, counts = _prepare_rows(session, rows)
        for index, (mapped, status, conflicts, target) in enumerate(prepared, 1):
            row = QinsiGoodsImportRow(
                import_batch_id=batch.id,
                source_file_name=mapped["source_file_name"],
                sheet_name=MASTER_SHEET,
                excel_row_number=index,
                qinsi_product_code=mapped.get("qinsi_goods_no"),
                barcode=mapped.get("jan"),
                parsed_data=_json_dumps({key: value for key, value in mapped.items() if key not in {"raw", "warnings", "errors"}}),
                raw_json=_json_dumps(mapped["raw"] | {
                    "_source_file": mapped["source_file_name"],
                    "_source_excel_row": mapped["excel_row_number"],
                }),
                validation_status=status,
                warnings=_json_dumps(mapped.get("warnings")) if mapped.get("warnings") else None,
                errors=_json_dumps(mapped.get("errors")) if mapped.get("errors") else None,
                conflict_json=_json_dumps(conflicts) if conflicts else None,
                product_id=target.id if target is not None else None,
            )
            session.add(row)
        statuses = [status for _, status, _, _ in prepared]
        batch.total_rows = len(prepared)
        batch.new_count = statuses.count("new")
        batch.update_count = statuses.count("update")
        batch.unchanged_count = statuses.count("unchanged")
        batch.skipped_count = 0
        batch.conflict_count = statuses.count("conflict")
        batch.error_count = statuses.count("error")
        batch.warning_count = sum(bool(mapped.get("warnings")) for mapped, _, _, _ in prepared)
        batch.success_count = counts["mark_imported_count"]
        batch.summary_json = _json_dumps({
            "parse_version": MASTER_PARSE_VERSION,
            "source": MASTER_SOURCE,
            "files": [{"filename": item.filename, "sha256": hashlib.sha256(item.content).hexdigest()} for item in files],
            **counts,
        })
        batch.status = "previewed"
        session.commit()
    except Exception as exc:
        session.rollback()
        raise ValueError(f"秦丝商品主数据解析失败：{exc}") from exc
    session.refresh(batch)
    return batch


def _decoded_row(row: QinsiGoodsImportRow) -> dict:
    mapped = json.loads(row.parsed_data)
    for key in MONEY_FIELDS | DECIMAL3_FIELDS:
        if mapped.get(key) is not None:
            mapped[key] = Decimal(str(mapped[key]))
    return mapped


def _refresh_batch_from_prepared(
    session: Session,
    batch: QinsiImportBatch,
    row_models: list[QinsiGoodsImportRow],
    prepared: list[tuple[dict, str, list[dict], Product | None]],
    counts: dict,
) -> None:
    statuses: list[str] = []
    for row, (mapped, status, conflicts, target) in zip(row_models, prepared, strict=True):
        statuses.append(status)
        row.qinsi_product_code = mapped.get("qinsi_goods_no")
        row.barcode = mapped.get("jan")
        row.parsed_data = _json_dumps({key: value for key, value in mapped.items() if key not in {"raw", "errors"}})
        row.validation_status = status
        row.warnings = _json_dumps(mapped.get("warnings")) if mapped.get("warnings") else None
        row.errors = _json_dumps(mapped.get("errors")) if mapped.get("errors") else None
        row.conflict_json = _json_dumps(conflicts) if conflicts else None
        row.product_id = target.id if target is not None else None
    batch.new_count = statuses.count("new")
    batch.update_count = statuses.count("update")
    batch.unchanged_count = statuses.count("unchanged")
    batch.skipped_count = statuses.count("skipped")
    batch.conflict_count = statuses.count("conflict")
    batch.error_count = statuses.count("error")
    batch.warning_count = sum(bool(mapped.get("warnings")) for mapped, _, _, _ in prepared)
    batch.success_count = counts["mark_imported_count"]
    summary = _json_loads(batch.summary_json, {})
    summary.update(counts)
    batch.summary_json = _json_dumps(summary)


def resolve_qinsi_master_conflict(
    session: Session,
    batch: QinsiImportBatch,
    row_id: int,
    *,
    resolution_type: str,
    action: str | None = None,
    note: str | None = None,
    auto_apply: bool = False,
    primary_product_id: int | None = None,
) -> QinsiConflictResolution:
    if batch.source_system != MASTER_SOURCE:
        raise ValueError("该任务不是秦丝商品主数据导入")
    if batch.status != "previewed":
        raise ValueError("只有 previewed 状态的秦丝导入可解除冲突")
    if resolution_type not in QINSI_CONFLICT_RESOLUTION_TYPES:
        raise ValueError("冲突类型无效")
    row_models = list(session.scalars(select(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
    ).order_by(QinsiGoodsImportRow.excel_row_number)))
    selected = next((item for item in row_models if item.id == row_id), None)
    if selected is None:
        raise LookupError("冲突行不存在")
    mapped_rows = [_decoded_row(item) for item in row_models]
    selected_index = row_models.index(selected)
    selected_mapped = mapped_rows[selected_index]
    conflicts = _json_loads(selected.conflict_json, [])
    if selected.validation_status != "conflict" and not conflicts:
        raise ValueError("该行当前不是待处理冲突")

    qinsi_codes = _qinsi_codes_for_barcode_group(mapped_rows, selected_mapped.get("jan")) or [selected_mapped.get("qinsi_goods_no")]
    conflict_key = qinsi_conflict_key(
        barcode=selected_mapped.get("jan"),
        qinsi_product_codes=qinsi_codes,
        conflict_type=resolution_type,
    )
    resolved_action = (action or QINSI_CONFLICT_RESOLUTION_TYPES[resolution_type]["action"]).strip()[:80]
    now = datetime.now(timezone.utc)
    resolution = session.scalar(
        select(QinsiConflictResolution).where(QinsiConflictResolution.conflict_key == conflict_key)
    )
    if resolution is None:
        resolution = QinsiConflictResolution(
            conflict_key=conflict_key,
            barcode=selected_mapped.get("jan"),
            qinsi_product_codes=_json_dumps(qinsi_codes),
            resolution_type=resolution_type,
            action=resolved_action,
            note=(note or "").strip() or None,
            auto_apply=auto_apply,
        )
        session.add(resolution)
        session.flush()
    else:
        resolution.barcode = selected_mapped.get("jan")
        resolution.qinsi_product_codes = _json_dumps(qinsi_codes)
        resolution.resolution_type = resolution_type
        resolution.action = resolved_action
        resolution.note = (note or "").strip() or None
        resolution.auto_apply = auto_apply
        resolution.updated_at = now
    if resolution_type == "true_duplicate":
        if primary_product_id is None:
            raise ValueError("同一商品重复登记需要先选择主商品")
        primary = session.get(Product, primary_product_id)
        if primary is None:
            raise ValueError("主商品不存在")
        if selected_mapped.get("jan"):
            merge_duplicate_jan_group(
                session,
                selected_mapped["jan"],
                primary_product_id=primary.id,
                actor="qinsi_conflict_resolution",
                commit=False,
            )

    if resolution_type in NON_MERGE_RESOLUTION_TYPES:
        for mapped in mapped_rows:
            key = qinsi_conflict_key(
                barcode=mapped.get("jan"),
                qinsi_product_codes=_qinsi_codes_for_barcode_group(mapped_rows, mapped.get("jan")) or [mapped.get("qinsi_goods_no")],
                conflict_type=resolution_type,
            )
            if key == conflict_key:
                _annotate_resolution(mapped, resolution, manual=True)
    prepared, counts = _prepare_rows(session, mapped_rows)
    _refresh_batch_from_prepared(session, batch, row_models, prepared, counts)
    session.commit()
    session.refresh(resolution)
    return resolution


def confirm_qinsi_master_import(
    session: Session,
    batch: QinsiImportBatch,
    *,
    allow_importing: bool = False,
) -> QinsiImportBatch:
    if batch.source_system != MASTER_SOURCE:
        raise ValueError("该任务不是秦丝商品主数据导入")
    if batch.status in {"completed", "completed_with_issues"}:
        return batch
    allowed = {"previewed", "importing"} if allow_importing else {"previewed"}
    if batch.status not in allowed:
        raise ValueError("导入任务状态不允许确认")
    rows = list(session.scalars(select(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
    ).order_by(QinsiGoodsImportRow.excel_row_number)))
    new_count = update_count = unchanged_count = 0
    now = datetime.now(timezone.utc)
    try:
        for row in rows:
            if row.validation_status in {"conflict", "error", "skipped"}:
                continue
            mapped = _decoded_row(row)
            prepared, _ = _prepare_rows(session, [mapped])
            _, status, conflicts, target = prepared[0]
            if conflicts or status in {"conflict", "error"}:
                row.validation_status = "conflict" if conflicts else "error"
                row.conflict_json = _json_dumps(conflicts) if conflicts else None
                continue
            if target is None:
                product_jan = None if mapped.get("force_no_product_jan") else mapped.get("jan")
                target = Product(
                    jan=product_jan,
                    name_cn=mapped.get("qinsi_name"),
                    name_source="qinsi_import",
                    display_name=format_product_display_name(mapped.get("qinsi_name"), None),
                    source="qinsi_import",
                    product_origin="qinsi",
                    status="qinsi_product_imported",
                )
                session.add(target)
                session.flush()
                action = "imported_new"
            else:
                action = "imported_updated" if _changed_product_fields(target, mapped) else "unchanged"
            if mapped.get("jan") and not mapped.get("force_no_product_jan") and target is not None and target.jan != mapped.get("jan"):
                merge_products_with_jan_into_primary(
                    session,
                    mapped["jan"],
                    target,
                    actor="qinsi_master_import",
                    commit=False,
                )
            _apply_product_fields(session, target, mapped, now)
            if mapped.get("jan") and not mapped.get("force_no_product_jan"):
                merge_duplicate_jan_group(
                    session,
                    mapped["jan"],
                    primary_product_id=target.id,
                    actor="qinsi_master_import",
                    commit=False,
                )
            row.product_id = target.id
            row.validation_status = action
            if action == "imported_new":
                new_count += 1
            elif action == "imported_updated":
                update_count += 1
            else:
                unchanged_count += 1
        batch.new_count = new_count
        batch.update_count = update_count
        batch.unchanged_count = unchanged_count
        batch.conflict_count = sum(row.validation_status == "conflict" for row in rows)
        batch.error_count = sum(row.validation_status == "error" for row in rows)
        batch.success_count = new_count + update_count + unchanged_count
        batch.status = "completed_with_issues" if batch.conflict_count or batch.error_count else "completed"
        batch.confirmed_at = now
        session.commit()
    except Exception:
        session.rollback()
        raise
    session.refresh(batch)
    return batch


def queue_qinsi_master_confirmation(session: Session, batch: QinsiImportBatch) -> None:
    if batch.status != "previewed":
        raise ValueError("导入任务状态不允许确认")
    batch.status = "importing"
    session.commit()


def process_queued_qinsi_master_confirmation(database_url: str, batch_id: int) -> None:
    engine = build_engine(database_url)
    try:
        with Session(engine) as session:
            batch = session.get(QinsiImportBatch, batch_id)
            if batch is not None:
                try:
                    confirm_qinsi_master_import(session, batch, allow_importing=True)
                except Exception as exc:
                    session.rollback()
                    failed = session.get(QinsiImportBatch, batch_id)
                    if failed is not None:
                        failed.status = "failed"
                        failed.error_message = str(exc)
                        session.commit()
    finally:
        engine.dispose()


def export_qinsi_master_audit_workbook(session: Session, batch: QinsiImportBatch) -> bytes:
    rows = list(session.scalars(select(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
    ).order_by(QinsiGoodsImportRow.excel_row_number)))
    workbook = Workbook()
    conflicts_sheet = workbook.active
    conflicts_sheet.title = "冲突明细"
    conflict_headers = [
        "source_file", "source_row_no", "秦丝商品名称", "秦丝货号", "秦丝商品条码", "秦丝单品条码",
        "按规则判定JAN", "冲突类型", "冲突原因", "本地命中的 Product.id", "本地 Product.jan",
        "本地商品名", "本地 qinsi_goods_no", "本地 qinsi_product_barcode", "本地 qinsi_unit_barcode",
        "本地状态", "建议处理", "是否可自动处理",
    ]
    error_headers = ["原始秦丝行", "错误原因", "发生阶段", "建议处理"]
    summary_headers = ["统计项", "数量"]
    conflicts_sheet.append(conflict_headers)
    errors_sheet = workbook.create_sheet("错误明细")
    errors_sheet.append(error_headers)
    summary_sheet = workbook.create_sheet("汇总")
    summary_sheet.append(summary_headers)
    conflict_counter: Counter[str] = Counter()
    conflict_total = error_total = 0
    for row in rows:
        if row.validation_status not in {"conflict", "error"}:
            continue
        mapped = _decoded_row(row)
        if row.validation_status == "conflict":
            conflict_total += 1
            conflicts = json.loads(row.conflict_json or "[]") or [{"field": "其他", "message": "身份冲突"}]
            conflict_type = _audit_conflict_type(conflicts)
            conflict_counter[conflict_type] += 1
            products = _audit_candidate_products(session, mapped, conflicts)
            conflicts_sheet.append([
                mapped.get("source_file_name"),
                mapped.get("excel_row_number"),
                mapped.get("qinsi_name"),
                mapped.get("qinsi_goods_no"),
                mapped.get("qinsi_product_barcode"),
                mapped.get("qinsi_unit_barcode"),
                mapped.get("jan"),
                conflict_type,
                _audit_messages(conflicts),
                _joined(product.id for product in products),
                _joined(product.jan for product in products),
                _joined(product.display_name or product.name_cn or product.name_ja for product in products),
                _joined(product.qinsi_product_code for product in products),
                _joined(product.qinsi_product_barcode for product in products),
                _joined(product.qinsi_unit_barcode for product in products),
                _joined(product.status for product in products),
                _audit_suggestion(conflict_type),
                "否",
            ])
        elif row.validation_status == "error":
            error_total += 1
            errors = json.loads(row.errors or "[]") or ["解析错误"]
            raw = json.loads(row.raw_json or "{}")
            errors_sheet.append([
                _json_dumps(raw),
                "；".join(errors),
                "解析/预览",
                "补齐或修正秦丝源行后重新导出并人工处理",
            ])
    summary_rows = [
        ("冲突总数", conflict_total),
        ("错误总数", error_total),
        ("同JAN多本地Product", conflict_counter["同一个判定JAN命中多个本地Product"]),
        ("同JAN多秦丝商品", conflict_counter["同一个有效JAN出现在多个秦丝商品"]),
        ("qinsi_goods_no冲突", conflict_counter["qinsi_goods_no对应多个本地Product"]),
        ("其他身份冲突", conflict_counter["秦丝行与现有Product身份明显不一致"]),
        ("其他", conflict_counter["其他"]),
    ]
    for item in summary_rows:
        summary_sheet.append(item)
    _format_audit_workbook(workbook)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _joined(values) -> str:
    return " | ".join(str(value) for value in values if value not in {None, ""})


def _audit_messages(conflicts: list[dict]) -> str:
    return "；".join(dict.fromkeys(str(conflict.get("message") or "身份冲突") for conflict in conflicts))


def _audit_conflict_type(conflicts: list[dict]) -> str:
    messages = _audit_messages(conflicts)
    if "同一最终JAN对应多个现有Product" in messages:
        return "同一个判定JAN命中多个本地Product"
    if "秦丝多个商品命中同一有效JAN" in messages:
        return "同一个有效JAN出现在多个秦丝商品"
    if "同一qinsi_goods_no对应多个本地Product" in messages:
        return "qinsi_goods_no对应多个本地Product"
    if (
        "货号映射商品与最终JAN命中商品不同" in messages
        or "最终JAN未命中同一商品" in messages
        or "货号命中商品已有不同JAN" in messages
    ):
        return "秦丝行与现有Product身份明显不一致"
    return "其他"


def _audit_suggestion(conflict_type: str) -> str:
    suggestions = {
        "同一个判定JAN命中多个本地Product": "人工确认哪个本地Product保留该JAN；本轮不自动合并",
        "同一个有效JAN出现在多个秦丝商品": "人工确认是否秦丝重复商品、套装/变体或条码填错；本轮不自动合并",
        "qinsi_goods_no对应多个本地Product": "人工确认秦丝货号唯一归属；本轮不自动改映射",
        "秦丝行与现有Product身份明显不一致": "人工比对JAN、货号、图片和商品名后决定保留/拆分/修正",
        "其他": "人工复核冲突原因后处理",
    }
    return suggestions.get(conflict_type, suggestions["其他"])


def _audit_candidate_products(session: Session, mapped: dict, conflicts: list[dict]) -> list[Product]:
    product_ids: set[int] = set()
    for conflict in conflicts:
        _collect_product_ids(conflict.get("existing_value"), product_ids)
    jan = mapped.get("jan")
    goods_no = mapped.get("qinsi_goods_no")
    if jan:
        product_ids.update(session.scalars(select(Product.id).where(Product.jan == jan)).all())
    if goods_no:
        product_ids.update(session.scalars(select(Product.id).where(Product.qinsi_product_code == goods_no)).all())
        product_ids.update(session.scalars(
            select(QinsiProductMapping.product_id).where(QinsiProductMapping.qinsi_product_code == goods_no)
        ).all())
    if not product_ids:
        return []
    return list(session.scalars(select(Product).where(Product.id.in_(product_ids)).order_by(Product.id)))


def _collect_product_ids(value, product_ids: set[int]) -> None:
    if isinstance(value, dict):
        if value.get("product_id") is not None:
            product_ids.add(int(value["product_id"]))
        for child in value.values():
            _collect_product_ids(child, product_ids)
    elif isinstance(value, list):
        for child in value:
            _collect_product_ids(child, product_ids)


def _format_audit_workbook(workbook: Workbook) -> None:
    identifier_headers = {
        "秦丝货号", "秦丝商品条码", "秦丝单品条码", "按规则判定JAN", "本地 Product.jan",
        "本地 qinsi_goods_no", "本地 qinsi_product_barcode", "本地 qinsi_unit_barcode",
    }
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        header_by_column = {cell.column: str(cell.value or "") for cell in sheet[1]}
        for column_cells in sheet.columns:
            header = header_by_column.get(column_cells[0].column, "")
            max_length = min(60, max(len(str(cell.value or "")) for cell in column_cells) + 2)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(12, max_length)
            if header in identifier_headers:
                for cell in column_cells[1:]:
                    cell.number_format = "@"


def qinsi_master_preview_statistics(session: Session, batch: QinsiImportBatch) -> dict:
    summary = json.loads(batch.summary_json or "{}")
    imported_new = session.scalar(select(func.count()).select_from(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
        QinsiGoodsImportRow.validation_status == "imported_new",
    )) or 0
    imported_updated = session.scalar(select(func.count()).select_from(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
        QinsiGoodsImportRow.validation_status == "imported_updated",
    )) or 0
    return {
        **summary,
        "imported_new": imported_new,
        "imported_updated": imported_updated,
    }


def load_master_input_files(paths: list[Path]) -> list[MasterInputFile]:
    return [MasterInputFile(filename=path.name, content=path.read_bytes()) for path in paths]
