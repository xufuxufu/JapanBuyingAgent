"""QinSi 报表 -> 进销存汇总 import (Phase 9A).

Establishes trustworthy sales/purchase FACTS for a user-declared date
range. Deliberately does not compute a recommended reorder quantity --
that is Phase 9B. QinsiInventorySnapshot stays the sole inventory
authority; this module never writes to it or to procurement execution
records.

Identity matching reuses the exact code/JAN cross-validation already
implemented for the QinSi inventory snapshot (app.qinsi_inventory._match_product)
and the project's one JAN validator (app.local_product.is_valid_jan) --
no second validator, no fuzzy name matching.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import secrets
import shutil
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.config import SALES_SUMMARY_PREVIEW_TEMP_DIR
from app.local_product import is_valid_jan
from app.models import Product, QinsiSalesSummaryLine, QinsiSalesSummarySnapshot
from app.qinsi_inventory import (
    FileRangeInfo,
    MultiFileCompletenessReport,
    _match_product,
    _parse_filename_range,
    range_completeness_from_file_infos,
)

EXPECTED_HEADERS = (
    "商品名称", "货号", "单品条码", "型号规格", "图片", "分类", "品牌",
    "采购量", "采购金额", "销售量", "销售金额", "购买客户数",
    "当前库存", "支撑销售天数", "仓库所属门店", "仓库",
)
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
ALLOWED_EXTENSIONS = (".xlsx",)


class SalesSummaryImportError(ValueError):
    pass


class PreviewTokenError(SalesSummaryImportError):
    """Preview token missing/expired/corrupted -- always a friendly,
    user-facing message (never a raw path or internal detail)."""


def _cell_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_quantity_or_count(value) -> tuple[int | None, str | None]:
    """"1" / "1个" / "12件" -> int. "-"/空 -> None (no anomaly). Garbage ->
    (None, error_message) -- never guessed. Returns (parsed, error)."""
    text = _cell_text(value)
    if text is None or text == "-":
        return None, None
    stripped = text
    for suffix in ("个", "件", "只", "支", "瓶", "盒", "台"):
        if stripped.endswith(suffix):
            stripped = stripped[: -len(suffix)]
            break
    stripped = stripped.replace(",", "").strip()
    try:
        number = Decimal(stripped)
    except InvalidOperation:
        return None, f"数量字段无法解析：{text!r}"
    if not number.is_finite() or number != number.to_integral_value():
        return None, f"数量字段不是整数：{text!r}"
    return int(number), None


def parse_support_days(value) -> tuple[int | None, str | None]:
    """支撑销售天数 (days of inventory support at the current sales rate).
    QinSi reports "∞" for rows with inventory but zero sales in the period
    (inventory / 0 = infinite support) -- a legitimate, well-understood
    sentinel, not an anomaly. Represented as None (same as "-"/empty),
    never as a fabricated large number."""
    text = _cell_text(value)
    if text is None or text in ("-", "∞", "inf", "INF", "无限"):
        return None, None
    return parse_quantity_or_count(value)


def parse_amount(value) -> tuple[Decimal | None, str | None]:
    text = _cell_text(value)
    if text is None or text == "-":
        return None, None
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None, f"金额字段无法解析：{text!r}"
    if not number.is_finite():
        return None, f"金额字段无法解析：{text!r}"
    return number, None


@dataclass(frozen=True, slots=True)
class ParsedSalesSummaryRow:
    row_no: int
    product_name: str | None
    qinsi_product_code: str | None
    barcode_raw: str | None
    jan_candidate: str | None
    purchase_quantity: int | None
    purchase_amount: Decimal | None
    sales_quantity: int | None
    sales_amount: Decimal | None
    customer_count: int | None
    reported_current_inventory: int | None
    reported_support_sales_days: int | None
    raw: dict
    parse_errors: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ParsedSalesSummaryWorkbook:
    headers: tuple[str, ...]
    rows: tuple[ParsedSalesSummaryRow, ...]


def parse_sales_summary_workbook(content: bytes) -> ParsedSalesSummaryWorkbook:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise SalesSummaryImportError("Excel工作簿无法解析") from exc
    try:
        sheet = workbook.active
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            raise SalesSummaryImportError("Excel没有表头行")
        headers = tuple(_cell_text(cell) or "" for cell in header_row)
        if headers[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
            raise SalesSummaryImportError(
                f"表头与秦丝『进销存汇总』导出格式不一致，期望：{EXPECTED_HEADERS}，实际：{headers}"
            )

        rows: list[ParsedSalesSummaryRow] = []
        for row_no, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if values is None or all(v is None for v in values):
                continue
            raw = dict(zip(headers, values))
            errors: list[str] = []

            name = _cell_text(raw.get("商品名称"))
            code = _cell_text(raw.get("货号"))
            barcode_raw = _cell_text(raw.get("单品条码"))
            jan_candidate = barcode_raw if barcode_raw and is_valid_jan(barcode_raw) else None

            purchase_quantity, err = parse_quantity_or_count(raw.get("采购量"))
            if err:
                errors.append(err)
            purchase_amount, err = parse_amount(raw.get("采购金额"))
            if err:
                errors.append(err)
            sales_quantity, err = parse_quantity_or_count(raw.get("销售量"))
            if err:
                errors.append(err)
            sales_amount, err = parse_amount(raw.get("销售金额"))
            if err:
                errors.append(err)
            customer_count, err = parse_quantity_or_count(raw.get("购买客户数"))
            if err:
                errors.append(err)
            reported_current_inventory, err = parse_quantity_or_count(raw.get("当前库存"))
            if err:
                errors.append(err)
            reported_support_sales_days, err = parse_support_days(raw.get("支撑销售天数"))
            if err:
                errors.append(err)

            rows.append(ParsedSalesSummaryRow(
                row_no=row_no, product_name=name, qinsi_product_code=code, barcode_raw=barcode_raw,
                jan_candidate=jan_candidate, purchase_quantity=purchase_quantity, purchase_amount=purchase_amount,
                sales_quantity=sales_quantity, sales_amount=sales_amount, customer_count=customer_count,
                reported_current_inventory=reported_current_inventory,
                reported_support_sales_days=reported_support_sales_days,
                raw={k: (str(v) if v is not None else None) for k, v in raw.items()},
                parse_errors=tuple(errors),
            ))
        return ParsedSalesSummaryWorkbook(headers=headers, rows=tuple(rows))
    finally:
        workbook.close()


def _validate_upload_files(files: list[tuple[str, bytes]]) -> None:
    if not files:
        raise SalesSummaryImportError("至少需要一个Excel文件")
    for filename, content in files:
        extension = Path(filename or "").suffix.lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise SalesSummaryImportError(f"只允许上传：{', '.join(ALLOWED_EXTENSIONS)}（{filename or '未命名文件'}）")
        if not content or len(content) > MAX_UPLOAD_BYTES:
            raise SalesSummaryImportError(f"Excel文件为空或超过允许大小（{filename or '未命名文件'}）")


def _combined_file_hash(files: list[tuple[str, bytes]]) -> str:
    if len(files) == 1:
        return hashlib.sha256(files[0][1]).hexdigest()
    digest = hashlib.sha256()
    for name, content in sorted(files, key=lambda item: item[0]):
        digest.update(f"{Path(name).name}:{len(content)}:".encode("utf-8"))
        digest.update(hashlib.sha256(content).digest())
    return digest.hexdigest()


def analyze_sales_summary_completeness(files: list[tuple[str, bytes]]) -> MultiFileCompletenessReport:
    """Filename-range + header-consistency check for a multi-file sales
    summary upload. Delegates the actual gap/overlap math to the same
    core the inventory snapshot flow uses."""
    infos: list[FileRangeInfo] = []
    headers_seen: set[tuple[str, ...]] = set()
    for filename, content in files:
        parsed = parse_sales_summary_workbook(content)
        headers_seen.add(parsed.headers)
        row_range = _parse_filename_range(filename)
        infos.append(FileRangeInfo(filename=Path(filename).name, range=row_range, row_count=len(parsed.rows)))
    return range_completeness_from_file_infos(infos, header_mismatch=len(headers_seen) > 1)


@dataclass(frozen=True, slots=True)
class _RowMatchResult:
    row: ParsedSalesSummaryRow
    source_file: str
    product: Product | None
    matching_method: str | None
    match_status: str


def _match_rows(session: Session, files: list[tuple[str, bytes]]) -> list[_RowMatchResult]:
    results: list[_RowMatchResult] = []
    for filename, content in files:
        parsed = parse_sales_summary_workbook(content)
        for row in parsed.rows:
            product, method, status, _conflict_pair = _match_product(
                session, code=row.qinsi_product_code, jan=row.jan_candidate, internal_sku=None,
            )
            results.append(_RowMatchResult(
                row=row, source_file=Path(filename).name, product=product,
                matching_method=method, match_status=status,
            ))
    return results


@dataclass(frozen=True, slots=True)
class DuplicateCodeConflict:
    qinsi_product_code: str
    occurrences: tuple[tuple[str, int], ...]  # (source_file, row_no)


@dataclass(frozen=True, slots=True)
class SalesSummaryPreview:
    completeness: MultiFileCompletenessReport
    period_start: datetime
    period_end: datetime
    period_days: int
    total_rows: int
    matched_count: int
    unmatched_count: int
    conflict_count: int
    sales_positive_sku_count: int
    sales_quantity_total: int
    sales_amount_total: Decimal
    purchase_quantity_total: int
    duplicate_qinsi_codes: tuple[DuplicateCodeConflict, ...]


def _validate_period(period_start: datetime, period_end: datetime) -> int:
    if period_end < period_start:
        raise SalesSummaryImportError("统计结束日期不能早于开始日期")
    return (period_end.date() - period_start.date()).days + 1


# ---------------------------------------------------------------------------
# Preview -> confirm staging (server-side, token-addressed)
#
# The browser must never round-trip raw Excel bytes through a hidden form
# field: Starlette's multipart/urlencoded parser caps any non-file FIELD at
# 1MB (python-multipart's max_part_size) -- a real multi-file, ~5000-row
# sales-summary upload blows past that immediately, however the field is
# encoded. Files an <input type="file"> streams in do NOT hit this cap (they
# take a different, unbounded code path), which is exactly why the original
# upload step worked while resubmitting the same bytes back on confirm did
# not. The fix: stage uploads to disk under a random token on preview, and
# have confirm read them back by token -- the browser only ever carries the
# token (a few dozen bytes) between the two requests.
# ---------------------------------------------------------------------------

PREVIEW_TOKEN_TTL_SECONDS = 3600
_TOKEN_PATTERN = re.compile(r"^[0-9a-f]{32}$")


def _preview_token_dir(token: str) -> Path | None:
    """None for anything that isn't exactly a 32-hex-char token -- this is
    checked BEFORE the token ever touches a path, so no user-controlled
    string (a filename, a stray query value, "../..") can reach the
    filesystem as a directory name."""
    if not _TOKEN_PATTERN.fullmatch(token or ""):
        return None
    root = SALES_SUMMARY_PREVIEW_TEMP_DIR.resolve()
    candidate = (root / token).resolve()
    try:
        candidate.relative_to(root)
    except ValueError:
        return None
    return candidate


def create_preview_token(
    files: list[tuple[str, bytes]], *, period_start: datetime, period_end: datetime,
    preview_summary: dict | None = None,
) -> str:
    """Stage uploaded files server-side; returns the token the browser
    should carry (in a single small hidden field) to the confirm step."""
    SALES_SUMMARY_PREVIEW_TEMP_DIR.mkdir(parents=True, exist_ok=True)
    token = secrets.token_hex(16)
    token_dir = SALES_SUMMARY_PREVIEW_TEMP_DIR / token
    files_dir = token_dir / "files"
    files_dir.mkdir(parents=True)
    manifest_files = []
    for index, (filename, content) in enumerate(files):
        stored_name = f"{index}.xlsx"
        (files_dir / stored_name).write_bytes(content)
        manifest_files.append({
            "stored_name": stored_name,
            "original_filename": Path(filename).name,
            "sha256": hashlib.sha256(content).hexdigest(),
            "size": len(content),
        })
    manifest = {
        "token": token,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "period_start": period_start.isoformat(),
        "period_end": period_end.isoformat(),
        "period_days": _validate_period(period_start, period_end),
        "files": manifest_files,
        # Informational only -- confirm always recomputes from the staged
        # files themselves, it never trusts this (or anything the browser
        # sends back) for the actual write.
        "preview_summary": preview_summary or {},
    }
    (token_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return token


@dataclass(frozen=True, slots=True)
class PreviewTokenData:
    token: str
    period_start: datetime
    period_end: datetime
    files: tuple[tuple[str, bytes], ...]


def load_preview_token(token: str) -> PreviewTokenData:
    """Re-reads period + files from the server-side manifest/staged files --
    never from anything the client resubmits. Raises PreviewTokenError with
    a friendly message for every failure mode (missing/expired/corrupt)."""
    token_dir = _preview_token_dir(token)
    if token_dir is None or not token_dir.is_dir():
        raise PreviewTokenError("预览已失效，请重新上传")
    manifest_path = token_dir / "manifest.json"
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        created_at = datetime.fromisoformat(manifest["created_at"])
    except (OSError, ValueError, KeyError):
        raise PreviewTokenError("预览已失效，请重新上传")
    if (datetime.now(timezone.utc) - created_at).total_seconds() > PREVIEW_TOKEN_TTL_SECONDS:
        discard_preview_token(token)
        raise PreviewTokenError("预览已过期（超过1小时），请重新上传")
    files = []
    try:
        for entry in manifest["files"]:
            file_path = token_dir / "files" / entry["stored_name"]
            content = file_path.read_bytes()
            if hashlib.sha256(content).hexdigest() != entry["sha256"]:
                raise PreviewTokenError("预览文件已损坏，请重新上传")
            files.append((entry["original_filename"], content))
    except (OSError, KeyError):
        raise PreviewTokenError("预览文件缺失，请重新上传")
    return PreviewTokenData(
        token=token,
        period_start=datetime.fromisoformat(manifest["period_start"]),
        period_end=datetime.fromisoformat(manifest["period_end"]),
        files=tuple(files),
    )


def discard_preview_token(token: str) -> None:
    """Best-effort cleanup; a missing/already-removed token is not an error."""
    token_dir = _preview_token_dir(token)
    if token_dir is not None and token_dir.is_dir():
        shutil.rmtree(token_dir, ignore_errors=True)


def preview_sales_summary_files(
    session: Session, files: list[tuple[str, bytes]], *, period_start: datetime, period_end: datetime,
) -> SalesSummaryPreview:
    """Read-only dry run: parses + matches every row but writes nothing."""
    _validate_upload_files(files)
    period_days = _validate_period(period_start, period_end)
    completeness = analyze_sales_summary_completeness(files)
    results = _match_rows(session, files)

    by_code: dict[str, list[tuple[str, int]]] = {}
    for result in results:
        if result.row.qinsi_product_code:
            by_code.setdefault(result.row.qinsi_product_code, []).append((result.source_file, result.row.row_no))
    duplicate_codes = tuple(
        DuplicateCodeConflict(qinsi_product_code=code, occurrences=tuple(occurrences))
        for code, occurrences in by_code.items() if len(occurrences) > 1
    )

    matched = sum(1 for r in results if r.match_status == "matched")
    unmatched = sum(1 for r in results if r.match_status == "unmatched")
    conflict = sum(1 for r in results if r.match_status == "conflict")
    sales_positive = sum(1 for r in results if (r.row.sales_quantity or 0) > 0)
    sales_qty_total = sum(r.row.sales_quantity or 0 for r in results)
    sales_amount_total = sum((r.row.sales_amount or Decimal(0)) for r in results) or Decimal(0)
    purchase_qty_total = sum(r.row.purchase_quantity or 0 for r in results)

    return SalesSummaryPreview(
        completeness=completeness, period_start=period_start, period_end=period_end, period_days=period_days,
        total_rows=len(results), matched_count=matched, unmatched_count=unmatched, conflict_count=conflict,
        sales_positive_sku_count=sales_positive, sales_quantity_total=sales_qty_total,
        sales_amount_total=sales_amount_total, purchase_quantity_total=purchase_qty_total,
        duplicate_qinsi_codes=duplicate_codes,
    )


def create_sales_summary_snapshot_from_files(
    session: Session, files: list[tuple[str, bytes]], *,
    period_start: datetime, period_end: datetime, now: datetime | None = None,
) -> tuple[QinsiSalesSummarySnapshot, bool]:
    """Create exactly one snapshot from one or more source files for one
    user-declared period. Never touches QinsiInventorySnapshot or
    procurement execution records."""
    _validate_upload_files(files)
    period_days = _validate_period(period_start, period_end)
    file_hash = _combined_file_hash(files)
    existing = session.scalar(
        select(QinsiSalesSummarySnapshot).where(QinsiSalesSummarySnapshot.file_hash == file_hash)
    )
    if existing is not None:
        return existing, True

    imported_at = now or datetime.now(timezone.utc)
    if len(files) == 1:
        original_filename = Path(files[0][0]).name[:255]
        file_content = files[0][1]
    else:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, content in files:
                archive.writestr(Path(name).name, content)
        file_content = buffer.getvalue()
        original_filename = f"秦丝进销存汇总_合并{len(files)}个文件_{imported_at:%Y%m%d}.zip"[:255]

    results = _match_rows(session, files)
    matched = sum(1 for r in results if r.match_status == "matched")
    unmatched = sum(1 for r in results if r.match_status == "unmatched")
    conflict = sum(1 for r in results if r.match_status == "conflict")

    snapshot = QinsiSalesSummarySnapshot(
        snapshot_no=f"QSS-{imported_at:%Y%m%d}-{uuid.uuid4().hex[:10].upper()}",
        period_start=period_start, period_end=period_end, period_days=period_days,
        imported_at=imported_at, original_filename=original_filename, file_hash=file_hash,
        file_content=file_content, total_rows=len(results), matched_rows=matched,
        unmatched_rows=unmatched, conflict_rows=conflict,
        status="completed" if not unmatched and not conflict else "completed_with_issues",
    )
    session.add(snapshot)
    session.flush()

    for result in results:
        error_parts = list(result.row.parse_errors)
        if result.match_status == "conflict":
            error_parts.append("货号与单品条码指向不同商品，未自动绑定")
        snapshot.lines.append(QinsiSalesSummaryLine(
            snapshot_id=snapshot.id, original_row_no=result.row.row_no,
            product_name_snapshot=result.row.product_name, qinsi_product_code=result.row.qinsi_product_code,
            jan_candidate=result.row.jan_candidate,
            product_id=result.product.id if result.product else None,
            match_status=result.match_status, matching_method=result.matching_method,
            purchase_quantity=result.row.purchase_quantity, purchase_amount=result.row.purchase_amount,
            sales_quantity=result.row.sales_quantity, sales_amount=result.row.sales_amount,
            customer_count=result.row.customer_count,
            reported_current_inventory=result.row.reported_current_inventory,
            reported_support_sales_days=result.row.reported_support_sales_days,
            raw_row_json=str(result.row.raw), error_message="；".join(error_parts) or None,
        ))

    try:
        session.commit()
    except IntegrityError:
        session.rollback()
        existing = session.scalar(
            select(QinsiSalesSummarySnapshot).where(QinsiSalesSummarySnapshot.file_hash == file_hash)
        )
        if existing is not None:
            return existing, True
        raise
    session.refresh(snapshot)
    return snapshot, False


def list_sales_summary_snapshots(session: Session) -> list[QinsiSalesSummarySnapshot]:
    return list(session.scalars(
        select(QinsiSalesSummarySnapshot).order_by(
            QinsiSalesSummarySnapshot.imported_at.desc(), QinsiSalesSummarySnapshot.id.desc(),
        )
    ))


def get_sales_summary_snapshot(session: Session, snapshot_id: int) -> QinsiSalesSummarySnapshot | None:
    return session.scalar(
        select(QinsiSalesSummarySnapshot)
        .where(QinsiSalesSummarySnapshot.id == snapshot_id)
        .options(selectinload(QinsiSalesSummarySnapshot.lines).selectinload(QinsiSalesSummaryLine.product))
    )


def get_sales_summary_snapshot_summary(session: Session, snapshot_id: int) -> QinsiSalesSummarySnapshot | None:
    """Snapshot row only, no eager line loading -- the header stats
    (total/matched/unmatched/conflict_rows) are precomputed columns set at
    import time, not derived from .lines, so the detail page's top summary
    never needs to touch the (potentially thousands-of-rows) line table."""
    return session.get(QinsiSalesSummarySnapshot, snapshot_id)


SALES_SUMMARY_LINE_PAGE_SIZES = (20, 50, 100)
DEFAULT_SALES_SUMMARY_LINE_PAGE_SIZE = 20


@dataclass(frozen=True, slots=True)
class SalesSummaryLinesPage:
    lines: list[QinsiSalesSummaryLine]
    total_count: int
    page: int
    page_size: int

    @property
    def total_pages(self) -> int:
        return max(1, -(-self.total_count // self.page_size))


def get_sales_summary_lines_page(
    session: Session, snapshot_id: int, *,
    match_status: str | None = None, search_query: str | None = None,
    page: int = 1, page_size: int = DEFAULT_SALES_SUMMARY_LINE_PAGE_SIZE,
) -> SalesSummaryLinesPage:
    """Real SQL-level pagination -- never loads the whole (potentially
    thousands-of-rows) line set to slice it in Python or hide rows with CSS.
    Filtering happens in the WHERE clause too, so total_count/total_pages
    reflect the filtered set, not the whole snapshot."""
    page_size = page_size if page_size in SALES_SUMMARY_LINE_PAGE_SIZES else DEFAULT_SALES_SUMMARY_LINE_PAGE_SIZE
    page = max(1, page)

    conditions = [QinsiSalesSummaryLine.snapshot_id == snapshot_id]
    if match_status and match_status != "all":
        conditions.append(QinsiSalesSummaryLine.match_status == match_status)
    if search_query:
        pattern = f"%{search_query}%"
        conditions.append(or_(
            QinsiSalesSummaryLine.product_name_snapshot.like(pattern),
            QinsiSalesSummaryLine.qinsi_product_code.like(pattern),
            QinsiSalesSummaryLine.jan_candidate.like(pattern),
        ))

    total_count = session.scalar(
        select(func.count()).select_from(QinsiSalesSummaryLine).where(*conditions)
    ) or 0
    lines = list(session.scalars(
        select(QinsiSalesSummaryLine).where(*conditions)
        .options(selectinload(QinsiSalesSummaryLine.product))
        .order_by(QinsiSalesSummaryLine.original_row_no)
        .limit(page_size).offset((page - 1) * page_size)
    ))
    return SalesSummaryLinesPage(lines=lines, total_count=total_count, page=page, page_size=page_size)


# ---------------------------------------------------------------------------
# Query services (Phase 9A: facts only, no recommended-quantity math)
# ---------------------------------------------------------------------------

@dataclass(frozen=True, slots=True)
class ProductSalesSummary:
    sales_quantity: int | None
    sales_known: bool
    sales_amount: Decimal | None
    customer_count: int | None
    data_start: datetime | None
    data_end: datetime | None


UNKNOWN_SALES_SUMMARY = ProductSalesSummary(
    sales_quantity=None, sales_known=False, sales_amount=None, customer_count=None,
    data_start=None, data_end=None,
)


def _latest_snapshot_for_period(session: Session, period_days: int) -> QinsiSalesSummarySnapshot | None:
    return session.scalar(
        select(QinsiSalesSummarySnapshot)
        .where(
            QinsiSalesSummarySnapshot.period_days == period_days,
            QinsiSalesSummarySnapshot.status.in_(("completed", "completed_with_issues")),
        )
        .order_by(QinsiSalesSummarySnapshot.imported_at.desc(), QinsiSalesSummarySnapshot.id.desc())
        .limit(1)
    )


def sales_summary_for_products(
    session: Session, product_ids: list[int], period_days: int,
) -> dict[int, ProductSalesSummary]:
    """Batched (no N+1) lookup of each product's sales fact for the given
    period length (e.g. 7 or 30), from the latest confirmed snapshot whose
    period_days matches exactly. A product with no row in that snapshot (or
    whose row didn't resolve to a Product) is UNKNOWN, never a silent zero.
    """
    unique_ids = list(dict.fromkeys(product_ids))
    if not unique_ids:
        return {}
    snapshot = _latest_snapshot_for_period(session, period_days)
    if snapshot is None:
        return {product_id: UNKNOWN_SALES_SUMMARY for product_id in unique_ids}

    rows = session.execute(
        select(
            QinsiSalesSummaryLine.product_id, QinsiSalesSummaryLine.sales_quantity,
            QinsiSalesSummaryLine.sales_amount, QinsiSalesSummaryLine.customer_count,
        ).where(
            QinsiSalesSummaryLine.snapshot_id == snapshot.id,
            QinsiSalesSummaryLine.product_id.in_(unique_ids),
            QinsiSalesSummaryLine.match_status == "matched",
        )
    ).all()
    found = {
        product_id: ProductSalesSummary(
            sales_quantity=sales_quantity or 0, sales_known=True, sales_amount=sales_amount,
            customer_count=customer_count, data_start=snapshot.period_start, data_end=snapshot.period_end,
        )
        for product_id, sales_quantity, sales_amount, customer_count in rows
    }
    return {product_id: found.get(product_id, UNKNOWN_SALES_SUMMARY) for product_id in unique_ids}


def product_replenishment_signals(session: Session, product_ids: list[int]) -> dict[int, dict]:
    """Lightweight fact-aggregation for Phase 9B to build on -- deliberately
    produces no recommended_quantity or any derived/predicted number, only
    facts already available elsewhere in the system."""
    from app.procurement_service import in_transit_quantity_for_products, reference_inventory_for_products
    from app.models import ProcurementDemandPlan, ProcurementPurchaseExecution

    unique_ids = list(dict.fromkeys(product_ids))
    if not unique_ids:
        return {}

    inventory_by_id = reference_inventory_for_products(session, unique_ids)
    in_transit_by_id = in_transit_quantity_for_products(session, unique_ids)
    sales_7d = sales_summary_for_products(session, unique_ids, 7)
    sales_30d = sales_summary_for_products(session, unique_ids, 30)

    planned_rows = session.execute(
        select(ProcurementDemandPlan.product_id, func.coalesce(func.sum(ProcurementDemandPlan.planned_quantity), 0))
        .where(ProcurementDemandPlan.product_id.in_(unique_ids), ProcurementDemandPlan.status == "planned")
        .group_by(ProcurementDemandPlan.product_id)
    ).all()
    planned_by_id = {product_id: int(total) for product_id, total in planned_rows}

    purchased_rows = session.execute(
        select(ProcurementDemandPlan.product_id, func.coalesce(func.sum(ProcurementPurchaseExecution.quantity), 0))
        .join(ProcurementPurchaseExecution, ProcurementPurchaseExecution.plan_id == ProcurementDemandPlan.id)
        .where(
            ProcurementDemandPlan.product_id.in_(unique_ids),
            ProcurementPurchaseExecution.status != "cancelled",
        )
        .group_by(ProcurementDemandPlan.product_id)
    ).all()
    purchased_by_id = {product_id: int(total) for product_id, total in purchased_rows}

    signals: dict[int, dict] = {}
    for product_id in unique_ids:
        inventory = inventory_by_id.get(product_id)
        summary_30d = sales_30d[product_id]
        summary_7d = sales_7d[product_id]
        signals[product_id] = {
            "china_inventory": inventory.china_quantity if inventory else None,
            "china_inventory_known": bool(inventory.china_known) if inventory else False,
            "japan_inventory": inventory.japan_quantity if inventory else None,
            "japan_inventory_known": bool(inventory.japan_known) if inventory else False,
            "sales_7d": summary_7d.sales_quantity,
            "sales_7d_known": summary_7d.sales_known,
            "sales_30d": summary_30d.sales_quantity,
            "sales_30d_known": summary_30d.sales_known,
            "in_transit_quantity": in_transit_by_id.get(product_id, 0),
            "planned_quantity": planned_by_id.get(product_id, 0),
            "purchased_quantity": purchased_by_id.get(product_id, 0),
            "sales_data_end": summary_30d.data_end or summary_7d.data_end,
        }
    return signals
