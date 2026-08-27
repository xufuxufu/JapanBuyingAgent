from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
import httpx
from openpyxl import load_workbook
from sqlalchemy import func, select

import app.qinsi_export as qinsi_export
from app.location_service import initialize_default_locations
from app.models import (
    PriceProviderAttempt, Product, PurchaseBatch, PurchaseBatchItem, QinsiExportJob, QinsiExportLine, QinsiPurchaseExportJob,
    QinsiPurchaseExportLine, QinsiPurchaseExportLineSource, Receipt, ReceiptBatch, ReceiptItem,
)
from app.qinsi_export import (
    QINSI_GOODS_TEMPLATE_HEADERS, QINSI_PURCHASE_TEMPLATE_HEADERS,
    cancel_qinsi_product_export_confirmation, confirm_qinsi_export,
    confirm_qinsi_product_export, create_qinsi_product_export,
    generate_purchase_batch_exports, retry_failed_qinsi_lines,
    qinsi_product_export_rows,
)
from app.schemas import QinsiExportConfirmationInput


def make_purchase(db, products: list[dict], *, same_warehouse: bool = False):
    locations = {location.display_name: location for location in initialize_default_locations(db)}
    models = [Product(
        name_cn=values["name"], name_ja=values.get("name_ja", "日本語名"), jan=values.get("jan"),
        qinsi_product_code=values.get("code"), product_origin=values.get("origin", "manual"),
        status=values.get("status", "qinsi_product_imported" if values.get("origin") == "qinsi" else "new_pending_review"),
        purchase_price=values.get("purchase_price", 120), sale_price=values.get("sale_price", 180),
        main_image_source_url=values.get("image_url"),
    ) for values in products]
    db.add_all(models)
    db.flush()
    gpt_batch = ReceiptBatch(batch_no=f"QINSI-EXPORT-{id(db)}-{len(models)}", status="confirmed", image_status="ready", gpt_status="reviewed")
    receipt = Receipt(
        batch=gpt_batch, raw_store_name="秦丝闭环测试店", purchased_at=datetime(2026, 7, 16, 3, 0, tzinfo=timezone.utc),
        raw_store_address="东京都测试区一丁目",
        confirmation_status="confirmed", review_status="reviewed", confirmed_at=datetime.now(timezone.utc),
    )
    db.add(receipt)
    db.flush()
    purchase = PurchaseBatch(
        batch_no=f"PB-QINSI-{receipt.id:08d}", receipt_id=receipt.id, gpt_batch_id=gpt_batch.id,
        purchased_at=receipt.purchased_at, store_name=receipt.raw_store_name, confirmed_at=receipt.confirmed_at,
        status="confirmed", default_initial_location_id=locations["日本家里库存"].id,
        default_qinsi_warehouse_id=locations["新日本仓库"].id,
    )
    db.add(purchase)
    db.flush()
    for index, product in enumerate(models, 1):
        receipt_item = ReceiptItem(
            receipt_id=receipt.id, line_no=index, raw_name=product.name_cn, jan_candidate=product.jan,
            product_id=product.id, match_status="matched_existing", quantity=index, unit_price=100 + index,
            discount_amount=0, line_total=(100 + index) * index, confidence=1, review_status="confirmed",
        )
        db.add(receipt_item)
        db.flush()
        if same_warehouse:
            warehouse = locations["新日本仓库"]
        elif not product.jan:
            warehouse = locations["无条码商品"]
        elif product.product_origin == "qinsi":
            warehouse = locations["新日本仓库"]
        else:
            warehouse = locations["日本家里库存"]
        db.add(PurchaseBatchItem(
            purchase_batch_id=purchase.id, product_id=product.id, receipt_item_id=receipt_item.id,
            quantity=index, unit_price=100 + index, discount_amount=0, actual_line_amount=(100 + index) * index,
            initial_location_id=locations["日本家里库存"].id, qinsi_target_warehouse_id=warehouse.id,
        ))
    db.commit()
    return purchase, models, locations


def workbook(content: bytes):
    return load_workbook(BytesIO(content), data_only=True)


def test_qinsi_new_product_export_uses_only_selected_products_and_detail_images(client):
    http, db, _ = client
    products = [
        Product(jan="4901234567894", name_cn="选择A", name_ja="選択A", status="new_pending_review", main_image_source_url="https://img.test/a.jpg"),
        Product(jan="4570110290418", name_cn="选择B", name_ja="選択B", status="new_pending_review", main_image_source_url="https://img.test/b.jpg"),
        Product(jan="00123457", name_cn="未选择C", name_ja="未選択C", status="new_pending_review"),
    ]
    db.add_all(products)
    db.commit()

    response = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": [str(products[0].id), str(products[1].id)]},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job = db.scalar(select(QinsiExportJob))
    assert db.scalar(select(func.count()).select_from(QinsiExportJob)) == 1
    assert db.scalar(select(func.count()).select_from(QinsiExportLine).where(QinsiExportLine.job_id == job.id)) == 2
    sheet = workbook(job.file_content)["商品导入"]
    exported_jans = {sheet.cell(row, 3).value for row in range(2, sheet.max_row + 1) if sheet.cell(row, 3).value}
    assert exported_jans == {products[0].jan, products[1].jan}
    detail = http.get(response.headers["location"])
    assert detail.status_code == 200
    assert "https://img.test/a.jpg" in detail.text and "https://img.test/b.jpg" in detail.text
    assert "暂无图片" not in detail.text


def test_html_qinsi_product_export_duplicates_redirect_or_render_links(client):
    http, db, _ = client
    same_a = Product(jan="4901417655387", name_cn="重复A", status="new_pending_review")
    same_b = Product(jan="4971710639322", name_cn="重复B", status="new_pending_review")
    other = Product(jan="4570110290418", name_cn="另一导出", status="new_pending_review")
    remaining = Product(jan="00123457", name_cn="剩余商品", status="new_pending_review")
    db.add_all([same_a, same_b, other, remaining])
    db.commit()
    same_job = create_qinsi_product_export(db, {same_a.id, same_b.id})
    other_job = create_qinsi_product_export(db, {other.id})

    same = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": [str(same_a.id), str(same_b.id)]},
        follow_redirects=False,
    )
    assert same.status_code == 303 and same.headers["location"] == f"/qinsi-product-exports/{same_job.id}"

    split = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": [str(same_a.id), str(other.id)]},
        follow_redirects=False,
    )
    assert split.status_code == 303 and split.headers["location"].startswith("/products?")
    split_page = http.get(split.headers["location"])
    assert split_page.status_code == 200
    assert f"/qinsi-product-exports/{same_job.id}" in split_page.text
    assert f"/qinsi-product-exports/{other_job.id}" in split_page.text
    assert '"detail"' not in split_page.text

    partial = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": [str(same_a.id), str(remaining.id)]},
        follow_redirects=False,
    )
    partial_page = http.get(partial.headers["location"])
    assert "仅生成剩余 1 件" in partial_page.text
    assert db.scalar(select(func.count()).select_from(QinsiExportJob)) == 2
    confirmed_remaining = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": str(remaining.id), "confirm_remaining": "1"},
        follow_redirects=False,
    )
    assert confirmed_remaining.status_code == 303
    assert db.scalar(select(func.count()).select_from(QinsiExportJob)) == 3


def test_new_existing_jan_no_jan_warehouse_template_and_full_source_tracking(client):
    _, db, _ = client
    receipt_named = Product(
        jan="4901234567894",
        name_cn="小票名称待确认",
        name_ja="票面简称",
        status="new_pending_completion",
        name_source="receipt",
        needs_review=True,
    )
    rich = Product(
        jan="4570110290418", name_cn="中文商品", name_ja="日本語商品", status="new_pending_review",
        purchase_price=Decimal("980"), main_image_source_url="https://img.test/first.jpg",
        local_image_path="data/products/qinsi-localized/local.jpg",
        display_image_url="/product-local-images/123?v=local",
    )
    db.add_all([receipt_named, rich])
    db.commit()
    job = create_qinsi_product_export(db, {receipt_named.id, rich.id})
    book = workbook(job.file_content)
    assert book.sheetnames == ["商品导入", "配置"]
    sheet = book["商品导入"]
    assert tuple(cell.value for cell in sheet[1]) == QINSI_GOODS_TEMPLATE_HEADERS
    rows = {sheet.cell(row, 3).value: row for row in range(2, 4)}
    missing_row, rich_row = rows[receipt_named.jan], rows[rich.jan]
    assert sheet.cell(missing_row, 1).value == "小票名称待确认|票面简称"
    assert sheet.cell(rich_row, 1).value == "中文商品|日本語商品"
    assert sheet.cell(rich_row, 2).value == rich.jan and sheet.cell(rich_row, 3).value == rich.jan
    assert sheet.cell(rich_row, 2).data_type == "s" and sheet.cell(rich_row, 3).data_type == "s"
    assert sheet.cell(rich_row, 7).value == "个"
    assert sheet.cell(rich_row, 8).value == 980
    assert sheet.cell(rich_row, 9).value == 980
    assert sheet.cell(rich_row, 8).value == sheet.cell(rich_row, 9).value
    assert sheet.cell(rich_row, 8).data_type == "n" and sheet.cell(rich_row, 9).data_type == "n"
    assert sheet.cell(missing_row, 8).value is None
    assert sheet.cell(missing_row, 9).value is None
    assert sheet.cell(missing_row, 10).value is None
    assert sheet.cell(rich_row, 10).value is None
    assert all(sheet.cell(row, column).value not in (0, "0", "0.00") for row in (missing_row, rich_row) for column in (8, 9, 10))
    assert sheet.cell(rich_row, 11).value == 100 and sheet.cell(rich_row, 12).value == "启用"
    assert sheet.cell(rich_row, 13).value == "启用" and sheet.cell(rich_row, 17).value == "停用"
    assert sheet.cell(rich_row, 19).value == rich.main_image_source_url and sheet.cell(rich_row, 24).value == "停用"
    assert "/product-local-images/" not in str(sheet.cell(rich_row, 19).value)
    assert "data/products" not in str(sheet.cell(rich_row, 19).value)
    assert all(sheet.cell(row, column).value is None for row in (missing_row, rich_row) for column in (26, 27, 28, 29))
    assert receipt_named.status == "new_pending_completion" and rich.status == "pending_qinsi_product_import"
    with pytest.raises(ValueError, match="已有待确认"):
        create_qinsi_product_export(db, {rich.id})


def test_qinsi_new_product_export_blocks_jan_without_usable_name(db_session):
    missing = Product(jan="4901234567894", name_cn="4901234567894", name_ja="缺商品", status="new_pending_completion")
    db_session.add(missing)
    db_session.commit()

    with pytest.raises(ValueError, match="缺少可用商品名"):
        create_qinsi_product_export(db_session, {missing.id})


def test_qinsi_new_product_export_accepts_single_language_names_without_price_or_image(db_session):
    only_ja = Product(jan="4901234567894", name_ja="日本語だけ商品", status="new_pending_completion")
    only_cn = Product(jan="4570110290418", name_cn="仅中文商品", status="new_pending_review")
    db_session.add_all([only_ja, only_cn])
    db_session.commit()

    job = create_qinsi_product_export(db_session, {only_ja.id, only_cn.id})
    sheet = workbook(job.file_content)["商品导入"]
    rows = {sheet.cell(row, 3).value: row for row in range(2, 4)}

    assert sheet.cell(rows[only_ja.jan], 1).value == "日本語だけ商品"
    assert sheet.cell(rows[only_cn.jan], 1).value == "仅中文商品"
    for row in rows.values():
        assert sheet.cell(row, 8).value is None
        assert sheet.cell(row, 9).value is None
        assert sheet.cell(row, 19).value is None
    assert only_ja.status == "new_pending_completion"
    assert only_cn.status == "pending_qinsi_product_import"


def test_qinsi_product_export_writes_yahoo_whitelist_image_and_blanks_non_whitelist(monkeypatch, db_session):
    yahoo_url = "https://item-shopping.c.yimg.jp/i/l/shop/4971710639322.jpg"
    rakuten_url = "https://thumbnail.image.rakuten.co.jp/@0_mall/shop/cabinet/4901417655387.jpg"
    official_url = "https://brand.example.com/images/4901417655387-1200.jpg"
    products = [
        Product(jan="4971710639322", name_cn="Yahoo图商品", status="new_pending_review", main_image_source_url=yahoo_url),
        Product(
            jan="4901417655387", name_cn="本地高清商品", status="new_pending_review",
            main_image_source_url=official_url, main_image_path="data/products/main/4901417655387.jpg",
            image_width=1200, image_height=1200, image_quality="normal", image_url=rakuten_url,
        ),
    ]
    db_session.add_all(products)
    db_session.commit()

    def fake_diagnostic(url):
        if url == yahoo_url:
            return qinsi_export.QinsiImageDiagnostic(
                url, url, 200, "image/jpeg", False, 800, 800, False, "write", "trusted_public_image",
            )
        if url == official_url:
            return qinsi_export.QinsiImageDiagnostic(
                url, None, None, None, False, None, None, False, "blank", "ConnectError",
            )
        return qinsi_export.QinsiImageDiagnostic(
            url, None, None, None, False, None, None, False, "blank", "private_or_tailscale_host",
        )

    monkeypatch.setattr(qinsi_export, "qinsi_image_diagnostic", fake_diagnostic)

    job = create_qinsi_product_export(db_session, {product.id for product in products})
    sheet = workbook(job.file_content)["商品导入"]
    rows = {sheet.cell(row, 3).value: row for row in range(2, 4)}

    assert sheet.cell(rows["4971710639322"], 19).value == yahoo_url
    assert sheet.cell(rows["4901417655387"], 19).value == official_url
    assert products[1].main_image_source_url == official_url
    assert products[1].image_width == 1200 and products[1].image_height == 1200


def test_qinsi_product_export_allows_local_image_without_public_url(db_session):
    product = Product(
        jan="4901234567894",
        name_cn="只有本地图商品",
        status="new_pending_review",
        main_image_path="data/products/main/4901234567894.jpg",
        display_image_url="/product-images/1",
    )
    db_session.add(product)
    db_session.commit()

    job = create_qinsi_product_export(db_session, {product.id})
    sheet = workbook(job.file_content)["商品导入"]
    assert sheet.cell(2, 19).value is None


def test_qinsi_product_export_manual_locked_image_does_not_use_old_auto_url(db_session):
    product = Product(
        jan="4901234567894",
        name_cn="人工包装图商品",
        status="new_pending_review",
        main_image_locked=True,
        main_image_path="data/products/main/manual.jpg",
        display_image_url="/product-images/1",
        main_image_source_url="https://images.qinsilk.com/old-auto.jpg",
        image_url="https://images.qinsilk.com/legacy-auto.jpg",
    )
    db_session.add(product)
    db_session.commit()

    job = create_qinsi_product_export(db_session, {product.id})
    sheet = workbook(job.file_content)["商品导入"]

    assert sheet.cell(2, 19).value is None


def test_qinsi_product_export_manual_locked_public_image_wins(db_session):
    public_manual = "https://images.qinsilk.com/manual-package.jpg"
    old_auto = "https://images.qinsilk.com/old-auto.jpg"
    product = Product(
        jan="4901234567894",
        name_cn="人工公网图商品",
        status="new_pending_review",
        main_image_locked=True,
        display_image_url=public_manual,
        main_image_source_url=old_auto,
    )
    db_session.add(product)
    db_session.commit()

    job = create_qinsi_product_export(db_session, {product.id})
    sheet = workbook(job.file_content)["商品导入"]

    assert sheet.cell(2, 19).value == public_manual


def test_qinsi_product_export_unlocked_uses_main_image_source_url(db_session):
    main_url = "https://images.qinsilk.com/main-source.jpg"
    legacy_url = "https://images.qinsilk.com/legacy.jpg"
    product = Product(
        jan="4901234567894",
        name_cn="未锁图商品",
        status="new_pending_review",
        main_image_source_url=main_url,
        image_url=legacy_url,
    )
    db_session.add(product)
    db_session.commit()

    job = create_qinsi_product_export(db_session, {product.id})
    sheet = workbook(job.file_content)["商品导入"]

    assert sheet.cell(2, 19).value == main_url


def test_qinsi_product_export_unlocked_keeps_main_then_image_order_and_blanks_private(monkeypatch, db_session):
    rakuten_490 = "https://thumbnail.image.rakuten.co.jp/@0_mall/shop/cabinet/4901417655387.jpg"
    rakuten_fallback = "https://thumbnail.image.rakuten.co.jp/@0_mall/shop/cabinet/4570110290418.jpg"
    tsnet_url = "https://xufu-cp.taile96adb.ts.net/images/private.jpg"
    products = [
        Product(
            jan="4901417655387", name_ja="テスト商品 1200ml", status="new_pending_review",
            main_image_source_url=rakuten_490, main_image_path="data/products/main/4901417655387.jpg",
            image_width=1200, image_height=1200, image_quality="normal",
        ),
        Product(jan="4570110290418", name_ja="楽天だけ商品", status="new_pending_review", main_image_source_url=rakuten_fallback),
        Product(jan="4901234567894", name_ja="私有图商品", status="new_pending_review", main_image_source_url=tsnet_url),
    ]
    db_session.add_all(products)
    db_session.commit()

    def fake_diagnostic(url):
        if "rakuten" in url:
            return qinsi_export.QinsiImageDiagnostic(
                url, url, 200, "image/jpeg", False, 1000, 1000, False, "blank", "host_not_qinsi_trusted",
            )
        if ".ts.net" in url:
            return qinsi_export.QinsiImageDiagnostic(
                url, None, None, None, False, None, None, False, "blank", "private_or_tailscale_host",
            )
        raise AssertionError(url)

    monkeypatch.setattr(qinsi_export, "qinsi_image_diagnostic", fake_diagnostic)

    job = create_qinsi_product_export(db_session, {product.id for product in products})
    sheet = workbook(job.file_content)["商品导入"]
    rows = {sheet.cell(row, 3).value: row for row in range(2, 5)}

    assert sheet.cell(rows["4901417655387"], 19).value == rakuten_490
    assert sheet.cell(rows["4570110290418"], 19).value == rakuten_fallback
    assert sheet.cell(rows["4901234567894"], 19).value is None
    assert products[0].main_image_source_url == rakuten_490
    assert products[0].qinsi_image_url is None


def test_qinsi_image_diagnostic_only_writes_trusted_public_decodable_non_thumbnail(monkeypatch):
    monkeypatch.setattr(
        qinsi_export.socket,
        "getaddrinfo",
        lambda host, port, type=None: [(None, None, None, None, ("8.8.8.8", 443))],
    )
    output = BytesIO()
    from PIL import Image

    Image.new("RGB", (800, 800), "white").save(output, "JPEG")

    class Client:
        def get(self, url, headers):
            request = httpx.Request("GET", url)
            return httpx.Response(
                200,
                content=output.getvalue(),
                headers={"content-type": "image/jpeg"},
                request=request,
            )

    trusted = qinsi_export.qinsi_image_diagnostic("https://images.qinsilk.com/a.jpg", client=Client())
    yahoo = qinsi_export.qinsi_image_diagnostic("https://item-shopping.c.yimg.jp/i/l/shop/a.jpg", client=Client())
    untrusted = qinsi_export.qinsi_image_diagnostic("https://thumbnail.image.rakuten.co.jp/a.jpg?_ex=128x128", client=Client())

    assert trusted and trusted.decision == "write" and trusted.width == 800
    assert yahoo and yahoo.decision == "write" and yahoo.reason == "trusted_public_image"
    assert untrusted and untrusted.decision == "blank" and untrusted.reason == "thumbnail"


def test_product_detail_edits_and_displays_structured_specs(client):
    http, db, _ = client
    product = Product(jan="4901234567894", name_ja="规格编辑商品", status="new_pending_completion")
    db.add(product)
    db.commit()

    response = http.post(
        f"/products/{product.id}",
        data={
            "name_cn": "",
            "name_ja": "规格编辑商品",
            "jan": product.jan,
            "main_image_source_url": "",
            "purchase_price": "",
            "status": "new_pending_completion",
            "specification": "W60×H60×D80mm / 140g / 100ml / 10个装",
            "net_weight_g": "140",
            "volume_ml": "100",
            "length_mm": "",
            "width_mm": "60",
            "height_mm": "60",
            "depth_mm": "80",
            "pack_quantity": "10",
            "spec_text": "W60×H60×D80mm / 140g / 100ml / 10个装",
            "actor": "测试",
            "reason": "规格编辑",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303
    db.refresh(product)
    assert product.net_weight_g == 140
    assert product.volume_ml == 100
    assert product.width_mm == 60 and product.height_mm == 60 and product.depth_mm == 80
    assert product.pack_quantity == 10

    detail = http.get(f"/products/{product.id}")
    assert detail.status_code == 200
    assert "140g / 100ml / W60×H60×D80mm / 10个装" in detail.text


def test_repeated_generation_and_download_reuse_same_records_and_bytes(client):
    _, db, _ = client
    purchase, products, _ = make_purchase(db, [
        {"name": "待导入商品", "jan": "4901234567894", "status": "new_pending_review"},
    ])
    with pytest.raises(ValueError, match="采购Excel已阻塞"):
        generate_purchase_batch_exports(db, purchase.id)
    product_job = create_qinsi_product_export(db, {products[0].id})
    confirm_qinsi_product_export(db, product_job, actor_name="测试操作人")
    assert products[0].status == "qinsi_product_imported" and products[0].qinsi_product_code is None
    purchase_job = generate_purchase_batch_exports(db, purchase.id)[0]
    assert purchase_job.export_type == "restock"
    cancel_qinsi_product_export_confirmation(db, product_job, actor_name="测试管理员")
    assert product_job.status == "exported" and product_job.cancelled_by == "测试管理员"
    assert product_job.cancelled_at is not None and products[0].status == "pending_qinsi_product_import"


def test_confirmed_product_export_repairs_product_status_and_products_queue(client):
    http, db, _ = client
    _, products, _ = make_purchase(db, [
        {"name": "历史已确认商品", "jan": "4901234567894", "status": "new_pending_review"},
    ])
    product_job = create_qinsi_product_export(db, {products[0].id})
    rows = qinsi_product_export_rows(db, product_job.id)
    for line, _ in rows:
        line.status = "confirmed"
    product_job.status = "confirmed"
    product_job.confirmed_at = datetime.now(timezone.utc)
    product_job.confirmed_by = "历史操作人"
    products[0].status = "pending_qinsi_product_import"
    db.commit()

    confirm_qinsi_product_export(db, product_job, actor_name="补同步")

    db.refresh(products[0])
    assert product_job.status == "confirmed"
    assert all(line.status == "confirmed" for line, _ in qinsi_product_export_rows(db, product_job.id))
    assert products[0].status == "qinsi_product_imported"
    page = http.get("/products?status=pending_qinsi_product_import")
    assert page.status_code == 200
    assert f'name="product_ids" value="{products[0].id}"' not in page.text


def test_products_queue_excludes_confirmed_import_even_before_status_repair(client):
    http, db, _ = client
    product = Product(name_cn="已确认但旧状态", name_ja="日本語", jan="4901234567894", status="pending_qinsi_product_import")
    db.add(product)
    db.flush()
    job = QinsiExportJob(status="confirmed", confirmed_at=datetime.now(timezone.utc), confirmed_by="历史操作人")
    db.add(job)
    db.flush()
    db.add(QinsiExportLine(job_id=job.id, product_id=product.id, qinsi_product_code=product.jan, product_name=product.name_cn, quantity=1, status="confirmed"))
    db.commit()

    page = http.get("/products?status=pending_qinsi_product_import")

    assert page.status_code == 200
    assert f'name="product_ids" value="{product.id}"' not in page.text


def test_all_success_is_idempotent_and_success_line_cannot_retry(db_session):
    purchase, products, locations = make_purchase(db_session, [
        {"name": "加权商品", "jan": "4901234567894", "origin": "qinsi"},
    ])
    first = purchase.items[0]
    first.quantity, first.unit_price, first.actual_line_amount = 1, 101, 101
    receipt_item = ReceiptItem(
        receipt_id=purchase.receipt_id, line_no=2, raw_name=products[0].name_cn,
        jan_candidate=products[0].jan, product_id=products[0].id, match_status="matched_existing",
        quantity=3, unit_price=200, discount_amount=0, line_total=600,
        confidence=1, review_status="confirmed",
    )
    db_session.add(receipt_item)
    db_session.flush()
    db_session.add(PurchaseBatchItem(
        purchase_batch_id=purchase.id, product_id=products[0].id, receipt_item_id=receipt_item.id,
        quantity=3, unit_price=200, discount_amount=0, actual_line_amount=600,
        initial_location_id=locations["日本家里库存"].id,
        qinsi_target_warehouse_id=locations["新日本仓库"].id,
    ))
    db_session.commit()
    job = generate_purchase_batch_exports(db_session, purchase.id)[0]
    book = workbook(job.file_content)
    sheet = book["采购单商品导入"]
    assert book.sheetnames == ["采购单商品导入"] and sheet.max_column == 7
    assert tuple(cell.value for cell in sheet[1]) == QINSI_PURCHASE_TEMPLATE_HEADERS
    assert [sheet.cell(2, column).value for column in range(1, 8)] == [
        products[0].jan, products[0].jan, "个", 4, 175.25, None, purchase.batch_no[:20],
    ]
    assert sheet["A2"].data_type == "s" and sheet["B2"].data_type == "s"
    assert job.line_count == 1 and len(job.lines) == 2
    assert generate_purchase_batch_exports(db_session, purchase.id)[0].id == job.id


def test_qinsi_purchase_excel_uses_actual_weighted_purchase_unit_price(db_session):
    purchase, products, _ = make_purchase(db_session, [
        {"name": "十件行金额商品", "jan": "4901234567894", "origin": "qinsi"},
    ])
    detail = purchase.items[0]
    detail.quantity = 10
    detail.unit_price = None
    detail.actual_line_amount = 2970
    db_session.commit()

    job = generate_purchase_batch_exports(db_session, purchase.id)[0]
    sheet = workbook(job.file_content)["采购单商品导入"]

    assert [sheet.cell(2, column).value for column in range(1, 6)] == [
        products[0].jan, products[0].jan, "个", 10, 297,
    ]


def test_multiple_receipts_create_distinct_batches_and_merge_export_tracks_all_batches(client):
    http, db_session, _ = client
    first, products, locations = make_purchase(db_session, [
        {"name": "跨批次加权商品", "jan": "4901234567894", "origin": "qinsi"},
    ], same_warehouse=True)
    first.items[0].quantity = 2
    first.items[0].unit_price = 100
    first.items[0].actual_line_amount = 180
    gpt_batch = ReceiptBatch(
        batch_no=f"QINSI-MERGE-{first.id}", status="confirmed", image_status="ready", gpt_status="reviewed",
    )
    receipt = Receipt(
        batch=gpt_batch, raw_store_name="第二张小票", confirmation_status="confirmed", review_status="reviewed",
        confirmed_at=datetime.now(timezone.utc),
    )
    db_session.add(receipt)
    db_session.flush()
    second = PurchaseBatch(
        batch_no=f"PB-MERGE-{receipt.id:08d}", receipt_id=receipt.id, gpt_batch_id=gpt_batch.id,
        store_name=receipt.raw_store_name, confirmed_at=receipt.confirmed_at, status="confirmed",
        default_initial_location_id=locations["日本家里库存"].id,
        default_qinsi_warehouse_id=locations["新日本仓库"].id,
    )
    db_session.add(second)
    db_session.flush()
    receipt_item = ReceiptItem(
        receipt_id=receipt.id, line_no=1, raw_name=products[0].name_cn, jan_candidate=products[0].jan,
        product_id=products[0].id, match_status="matched_existing", quantity=3, unit_price=200,
        discount_amount=0, line_total=600, confidence=1, review_status="confirmed",
    )
    db_session.add(receipt_item)
    db_session.flush()
    db_session.add(PurchaseBatchItem(
        purchase_batch_id=second.id, product_id=products[0].id, receipt_item_id=receipt_item.id,
        quantity=3, unit_price=200, discount_amount=0, actual_line_amount=600,
        initial_location_id=locations["日本家里库存"].id,
        qinsi_target_warehouse_id=locations["新日本仓库"].id,
    ))
    db_session.commit()

    response = http.post(
        "/purchase-batches/qinsi-exports/merge",
        data={"purchase_batch_ids": [str(first.id), str(second.id)]},
        follow_redirects=False,
    )
    assert response.status_code == 303
    job = db_session.scalar(select(QinsiPurchaseExportJob))
    download = http.get(f"/qinsi-exports/{job.id}/download")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = workbook(download.content)["采购单商品导入"]
    assert job.selected_batch_ids == sorted([first.id, second.id])
    assert {line.purchase_batch_id for line in job.lines} == {first.id, second.id}
    assert [sheet.cell(2, column).value for column in range(1, 6)] == [
        products[0].jan, products[0].jan, "个", 5, 156,
    ]


def test_merge_purchase_batches_less_than_two_returns_html_error(client):
    http, db, _ = client
    first, _, _ = make_purchase(db, [
        {"name": "单选商品", "jan": "4901234567894", "origin": "qinsi"},
    ], same_warehouse=True)

    response = http.post(
        "/purchase-batches/qinsi-exports/merge",
        data={"purchase_batch_ids": [str(first.id)]},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "请至少选择2个待入库批次" in response.text


def test_merge_purchase_batches_blocking_products_return_html_details(client):
    http, db, _ = client
    first, products, locations = make_purchase(db, [
        {"name": "可入库商品", "jan": "4901234567894", "origin": "qinsi"},
    ], same_warehouse=True)
    blocker = Product(name_cn="阻塞新品", name_ja="未登録", jan="4570110290418", status="new_pending_review")
    db.add(blocker)
    gpt_batch = ReceiptBatch(
        batch_no=f"QINSI-BLOCK-{first.id}", status="confirmed", image_status="ready", gpt_status="reviewed",
    )
    receipt = Receipt(
        batch=gpt_batch, raw_store_name="阻塞小票", confirmation_status="confirmed", review_status="reviewed",
        confirmed_at=datetime.now(timezone.utc),
    )
    db.add(receipt)
    db.flush()
    second = PurchaseBatch(
        batch_no=f"PB-BLOCK-{receipt.id:08d}", receipt_id=receipt.id, gpt_batch_id=gpt_batch.id,
        store_name=receipt.raw_store_name, confirmed_at=receipt.confirmed_at, status="confirmed",
        default_initial_location_id=locations["日本家里库存"].id,
        default_qinsi_warehouse_id=locations["新日本仓库"].id,
    )
    db.add(second)
    db.flush()
    receipt_item = ReceiptItem(
        receipt_id=receipt.id, line_no=1, raw_name=blocker.name_cn, jan_candidate=blocker.jan,
        product_id=blocker.id, match_status="new_product", quantity=1, unit_price=200,
        discount_amount=0, line_total=200, confidence=1, review_status="confirmed",
    )
    db.add(receipt_item)
    db.flush()
    db.add(PurchaseBatchItem(
        purchase_batch_id=second.id, product_id=blocker.id, receipt_item_id=receipt_item.id,
        quantity=1, unit_price=200, discount_amount=0, actual_line_amount=200,
        initial_location_id=locations["日本家里库存"].id,
        qinsi_target_warehouse_id=locations["新日本仓库"].id,
    ))
    db.commit()

    response = http.post(
        "/purchase-batches/qinsi-exports/merge",
        data={"purchase_batch_ids": [str(first.id), str(second.id)]},
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert "存在未导入秦丝的新商品" in response.text and "阻塞新品" in response.text
    assert f'value="{first.id}" aria-label="选择 {first.batch_no}" checked' in response.text


def test_local_price_check_is_immediate_and_uses_confirmed_receipt_actual_unit_price(client):
    http, db, _ = client
    _, products, _ = make_purchase(db, [{
        "name": "本地查价商品", "name_ja": "ローカル商品", "jan": "4901234567894",
        "origin": "qinsi", "purchase_price": 999, "image_url": "https://img.test/local.jpg",
    }])
    response = http.post("/price-check", data={"jan": products[0].jan}, follow_redirects=False)
    assert response.status_code == 303
    result = http.get(response.headers["location"])
    assert result.status_code == 200
    for expected in (
        "本地已有商品", "本地查价商品|ローカル商品", "https://img.test/local.jpg",
        "历史最低采购价", "上一次采购价", "¥101", "秦丝闭环测试店", "东京都测试区一丁目",
    ):
        assert expected in result.text
    assert db.scalar(select(func.count()).select_from(PriceProviderAttempt)) == 0


def test_partial_failure_only_failed_line_reexports_and_success_stays_locked(db_session):
    purchase, _, _ = make_purchase(db_session, [
        {"name": "部分成功A", "jan": "4901234567894", "origin": "qinsi"},
        {"name": "部分失败B", "jan": "4570110290418", "origin": "qinsi"},
    ], same_warehouse=True)
    job = generate_purchase_batch_exports(db_session, purchase.id)[0]
    failed_line, success_line = job.lines
    result = confirm_qinsi_export(db_session, job, QinsiExportConfirmationInput(
        result="partial_failure", failed_line_ids={failed_line.id},
    ))
    failed_line, success_line = result.lines
    assert result.status == "partially_failed"
    assert failed_line.status == "failed" and failed_line.source.is_active is False
    assert success_line.status == "imported" and success_line.source.is_active is True
    with pytest.raises(ValueError, match="失败行"):
        retry_failed_qinsi_lines(db_session, result, {success_line.id})
    retry = retry_failed_qinsi_lines(db_session, result, {failed_line.id})
    repeated_retry = retry_failed_qinsi_lines(db_session, result, {failed_line.id})
    assert retry.parent_export_job_id == result.id and retry.line_count == 1
    assert repeated_retry.id == retry.id
    assert retry.lines[0].purchase_batch_item_id == failed_line.purchase_batch_item_id
    assert retry.lines[0].purchase_batch_item_id != success_line.purchase_batch_item_id


def test_all_failed_rows_can_be_selected_for_retry(db_session):
    purchase, _, _ = make_purchase(db_session, [
        {"name": "全部失败A", "jan": "4901234567894", "origin": "qinsi"},
        {"name": "全部失败B", "jan": "4570110290418", "origin": "qinsi"},
    ], same_warehouse=True)
    job = generate_purchase_batch_exports(db_session, purchase.id)[0]
    result = confirm_qinsi_export(db_session, job, QinsiExportConfirmationInput(result="all_failed"))
    assert result.status == "failed" and all(line.status == "failed" and not line.source.is_active for line in result.lines)
    retry = retry_failed_qinsi_lines(db_session, result, {line.id for line in result.lines})
    assert retry.status == "generated" and retry.line_count == 2


def test_qinsi_purchase_confirm_buttons_and_invalid_result_render_chinese_error(client):
    http, db, _ = client
    scenarios = (
        ("all_success", "imported", ["4901234567894"]),
        ("partial_failure", "partially_failed", ["00123457", "4901417655387"]),
        ("all_failed", "failed", ["4971710639322", "4903301142797", "4903301142995"]),
    )
    for result, expected_status, jans in scenarios:
        rows = [
            {"name": f"确认按钮{result}{index}", "jan": jan, "origin": "qinsi"}
            for index, jan in enumerate(jans, 1)
        ]
        purchase, _, _ = make_purchase(db, rows, same_warehouse=True)
        job = generate_purchase_batch_exports(db, purchase.id)[0]
        page = http.get(f"/qinsi-exports/{job.id}")
        assert page.status_code == 200
        assert 'name="result" id="confirm-result"' in page.text
        assert f'data-result="{result}"' in page.text
        data = {"result": result, "actor_name": "测试确认"}
        if result == "partial_failure":
            data["failed_line_ids"] = str(job.lines[0].id)
        response = http.post(f"/qinsi-exports/{job.id}/confirm", data=data, follow_redirects=False)
        assert response.status_code == 303
        db.refresh(job)
        assert job.status == expected_status

    purchase, _, _ = make_purchase(db, [
        {"name": "非法确认值", "jan": "0490123456789", "origin": "qinsi"},
        {"name": "非法确认值2", "jan": "4901301230633", "origin": "qinsi"},
        {"name": "非法确认值3", "jan": "4901301230640", "origin": "qinsi"},
        {"name": "非法确认值4", "jan": "4901301230657", "origin": "qinsi"},
    ], same_warehouse=True)
    job = generate_purchase_batch_exports(db, purchase.id)[0]
    response = http.post(f"/qinsi-exports/{job.id}/confirm", data={"result": "success"}, follow_redirects=True)
    assert response.status_code == 422
    assert "确认结果无效" in response.text
    assert "Input should be" not in response.text


def test_qinsi_purchase_confirm_list_shows_product_images(client):
    http, db, _ = client
    purchase, _, _ = make_purchase(db, [{
        "name": "确认列表图片", "jan": "4901234567894", "origin": "qinsi",
        "image_url": "https://img.test/qinsi-confirm.jpg",
    }], same_warehouse=True)
    job = generate_purchase_batch_exports(db, purchase.id)[0]

    response = http.get(f"/qinsi-exports/{job.id}")

    assert response.status_code == 200
    assert "<th>图片</th>" in response.text
    assert "https://img.test/qinsi-confirm.jpg" in response.text
    assert "暂无图片" in response.text


def test_receipt_source_link_opens_confirmed_review_at_item_anchor(client):
    http, db, _ = client
    purchase, _, _ = make_purchase(db, [{
        "name": "小票链接商品", "jan": "4901234567894", "origin": "qinsi",
    }], same_warehouse=True)
    item = purchase.items[0].receipt_item

    response = http.get(
        f"/receipts/{purchase.gpt_batch_id}/review?receipt_id={purchase.receipt_id}#receipt-item-{item.id}"
    )

    assert response.status_code == 200
    assert f'id="receipt-item-{item.id}"' in response.text
    assert "此批次已最终确认，当前为只读状态" in response.text


def test_new_pending_completion_with_receipt_name_and_purchase_price_exports_and_is_listed(client):
    http, db, _ = client
    purchase, products, _ = make_purchase(db, [{
        "name": "4901234567894", "name_ja": "缺商品", "jan": "4901234567894",
        "status": "new_pending_completion", "purchase_price": None, "sale_price": None,
    }])
    product = products[0]
    product.display_name = f"{product.jan}|缺商品"
    purchase.items[0].receipt_item.raw_name = "小票原始商品名"
    db.commit()

    blocked = http.get(f"/purchase-batches/{purchase.id}")
    assert blocked.status_code == 200
    assert "被 1 个新商品阻塞" in blocked.text
    assert f"qinsi_product_ids={product.id}" in blocked.text

    listing = http.get(f"/products?qinsi_product_ids={product.id}")
    assert listing.status_code == 200
    assert "未登录秦丝的新商品" in listing.text
    assert product.jan in listing.text
    assert f"/price-check?jan={product.jan}" in listing.text

    created = http.post("/products/qinsi-new-exports", data={"product_ids": str(product.id)}, follow_redirects=False)
    assert created.status_code == 303
    job = db.scalar(select(QinsiExportJob).where(QinsiExportJob.id.is_not(None)).order_by(QinsiExportJob.id.desc()))
    sheet = workbook(job.file_content)["商品导入"]
    assert sheet.cell(2, 1).value == "小票原始商品名"
    assert sheet.cell(2, 3).value == product.jan
    assert sheet.cell(2, 8).value == 101
    assert product.purchase_price == Decimal("101")
    assert purchase.items[0].unit_price == 101


def test_blocked_six_new_pending_completion_jans_are_all_visible_and_exportable(client):
    http, db, _ = client
    jans = [
        "4960919412621", "4960919412584", "4960919412577",
        "4960919412614", "8800366242630", "8800366242654",
    ]
    products = [
        Product(jan=jan, name_ja=f"票面商品{index}", status="new_pending_completion")
        for index, jan in enumerate(jans, 1)
    ]
    db.add_all(products)
    db.commit()
    ids = ",".join(str(product.id) for product in products)

    listing = http.get(f"/products?qinsi_product_ids={ids}")

    assert listing.status_code == 200
    for jan in jans:
        assert jan in listing.text
        assert f"/price-check?jan={jan}" in listing.text
    response = http.post(
        "/products/qinsi-new-exports",
        data={"product_ids": [str(product.id) for product in products]},
        follow_redirects=False,
    )
    assert response.status_code == 303


def test_pages_show_three_work_queues_mobile_selection_and_duplicate_submit_guard(client):
    http, db, _ = client
    purchase, products, _ = make_purchase(db, [
        {"name": "页面新品", "jan": "4901234567894", "status": "new_pending_review"},
    ])
    blocked = http.get(f"/purchase-batches/{purchase.id}")
    assert blocked.status_code == 200 and "被 1 个新商品阻塞" in blocked.text
    assert "打开未登录商品列表" in blocked.text
    products_page = http.get("/products?status=new_pending_review")
    assert "未登录秦丝的新商品" in products_page.text and "全选" in products_page.text
    assert "生成秦丝新商品Excel（0）" in products_page.text and "qinsi-product-export-submit" in products_page.text
    created = http.post(
        "/products/qinsi-new-exports", data={"product_ids": str(products[0].id)}, follow_redirects=False,
    )
    assert created.status_code == 303
    product_job = db.scalar(select(QinsiExportJob))
    detail = http.get(created.headers["location"])
    assert "秦丝商品导入成功" in detail.text and "采购入库分开确认" in detail.text
    http.post(f"/qinsi-product-exports/{product_job.id}/confirm", data={"actor_name": "页面操作人"})
    purchase_page = http.get(f"/purchase-batches/{purchase.id}")
    assert "生成秦丝采购单商品Excel" in purchase_page.text and "single-submit" in purchase_page.text
    response = http.post(f"/purchase-batches/{purchase.id}/qinsi-exports", follow_redirects=False)
    assert response.status_code == 303
    purchase_job = db.scalar(select(QinsiPurchaseExportJob))
    purchase_detail = http.get(f"/qinsi-exports/{purchase_job.id}")
    assert "秦丝采购入库成功" in purchase_detail.text and "新日本仓库" in purchase_detail.text
