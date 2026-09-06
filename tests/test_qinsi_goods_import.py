from __future__ import annotations

import json
from datetime import datetime, timezone
from io import BytesIO
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.location_service import initialize_default_locations
from app.models import (
    DurableBackgroundJob,
    Location,
    Product,
    ProductBarcode,
    PurchaseBatch,
    PurchaseBatchItem,
    QinsiGoodsImportRow,
    QinsiImportBatch,
    QinsiInventorySnapshot,
    QinsiInventorySnapshotLine,
    QinsiMasterValue,
    Receipt,
    ReceiptBatch,
    ReceiptItem,
)
from app.qinsi_goods_import import (
    _mapped_row,
    confirm_import,
    create_import_preview,
    parse_qinsi_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
REAL_BOOK = ROOT / "tests" / "testExcelFiles" / "goodsImportTemplate_1-2000.xlsx"
REFERENCE_BOOK = ROOT / "reference" / "qinsi" / "goodsImportTemplate已有商品模版-可到导入到本地数据库.xlsx"
requires_real_book = pytest.mark.skipif(
    not REAL_BOOK.exists(),
    reason="本地真实秦丝Excel fixture未提供（tests/testExcelFiles 已被 gitignore）",
)


@requires_real_book
def test_real_qinsi_workbook_parses_logical_values_and_dynamic_config():
    workbook = parse_qinsi_workbook(REAL_BOOK.read_bytes())
    assert len(workbook.rows) == 2000
    assert len(workbook.headers) == 30
    assert workbook.inventory_warehouse == "新日本仓库"
    assert {key: len(value) for key, value in workbook.config.items()} == {
        "brand": 84,
        "category": 53,
        "unit": 11,
        "warehouse": 5,
        "product_status": 2,
        "points_status": 2,
    }
    assert {
        "blandValidate", "categoryValidate", "unitValidate", "storehouseValidate",
        "onSaleValidate", "onClientPointEnable",
    } <= set(workbook.defined_names)

    raw_rows = [raw for _, raw in workbook.rows]
    assert all(tuple(raw) == workbook.headers for raw in raw_rows)
    assert all(len(raw) == 30 for raw in raw_rows)
    assert all(raw["产地"] is None and raw["适用年龄"] is None and raw["库位"] is None for raw in raw_rows)
    assert not any(value == "184" for raw in raw_rows for value in raw.values())

    mapped = [_mapped_row(raw, workbook.config)[0] for raw in raw_rows]
    codes = [row["qinsi_product_code"] for row in mapped if row["qinsi_product_code"]]
    barcodes = [row["jan"] for row in mapped if row["jan"]]
    assert len(codes) == len(set(codes)) == 2000
    assert len(barcodes) == 865
    assert sum(row["jan"] is None for row in mapped) == 1135
    assert sum(
        row["purchase_price"] is not None and row["purchase_price"] % 1 != 0
        for row in mapped
    ) == 85
    assert sum(
        row["sale_price"] is not None and row["sale_price"] % 1 != 0
        for row in mapped
    ) == 4

    counted = [row["counted_inventory_quantity"] for row in mapped if row["counted_inventory_quantity"] is not None]
    assert len(counted) == 626
    assert sum(value == 0 for value in counted) == 575
    assert sum(value > 0 for value in counted) == 51
    assert sum(value < 0 for value in counted) == 0


@requires_real_book
def test_real_qinsi_preview_confirm_and_file_hash_idempotency(db_session):
    initialize_default_locations(db_session)
    content = REAL_BOOK.read_bytes()
    batch = create_import_preview(
        db_session,
        REAL_BOOK.name,
        content,
        business_batch_key="QINSI-5412-LOCAL-ACCEPTANCE",
    )
    assert batch.status == "previewed"
    assert (
        batch.total_rows, batch.new_count, batch.update_count, batch.unchanged_count,
        batch.skipped_count, batch.conflict_count, batch.error_count, batch.warning_count,
    ) == (2000, 1974, 0, 0, 0, 24, 2, 0)

    import_rows = list(db_session.scalars(select(QinsiGoodsImportRow).where(
        QinsiGoodsImportRow.import_batch_id == batch.id,
    ).order_by(QinsiGoodsImportRow.excel_row_number)))
    assert len(import_rows) == 2000
    assert sum(row.barcode is not None for row in import_rows) == 865
    assert sum(row.barcode is None for row in import_rows) == 1135
    assert not any(row.barcode == "184" for row in import_rows)
    assert not any("可疑值184" in (row.warnings or "") for row in import_rows)
    assert not any("重复非空条码：184" in (row.conflict_json or "") for row in import_rows)
    assert sum(row.validation_status == "error" for row in import_rows) == 2
    assert sum(row.validation_status == "conflict" for row in import_rows) == 24
    assert all(
        "合法 JAN-8/JAN-13" in (row.errors or "")
        for row in import_rows if row.validation_status == "error"
    )
    assert all(len(json.loads(row.raw_json)) == 30 for row in import_rows)

    confirm_import(db_session, batch)
    assert batch.status == "completed_with_issues"
    assert (batch.new_count, batch.update_count, batch.unchanged_count) == (1974, 0, 0)
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1974
    assert db_session.scalar(
        select(func.count()).select_from(Product).where(Product.product_note.is_not(None))
    ) == 0
    assert db_session.scalar(
        select(func.count()).select_from(Product).where(Product.name_ja.is_not(None))
    ) > 0
    assert sum(
        product.purchase_price is not None and product.purchase_price % Decimal("1") != 0
        for product in db_session.scalars(select(Product))
    ) == 85
    assert sum(
        product.sale_price is not None and product.sale_price % Decimal("1") != 0
        for product in db_session.scalars(select(Product))
    ) == 4

    master_counts = dict(db_session.execute(
        select(QinsiMasterValue.master_type, func.count())
        .group_by(QinsiMasterValue.master_type)
    ).all())
    assert master_counts == {
        "brand": 84,
        "category": 53,
        "unit": 11,
        "warehouse": 5,
        "product_status": 2,
        "points_status": 2,
    }

    snapshot = db_session.scalar(select(QinsiInventorySnapshot))
    assert snapshot is not None
    assert snapshot.source_system == "qinsi"
    assert snapshot.snapshot_type == "counted_inventory"
    assert snapshot.source_import_batch_id == batch.id
    assert snapshot.original_filename == REAL_BOOK.name
    assert (snapshot.total_rows, snapshot.success_rows, snapshot.unmatched_rows, snapshot.exception_rows) == (616, 616, 0, 0)
    lines = list(db_session.scalars(select(QinsiInventorySnapshotLine).where(
        QinsiInventorySnapshotLine.snapshot_id == snapshot.id,
    )))
    assert len(lines) == 616
    assert {line.raw_warehouse_name for line in lines} == {"新日本仓库"}
    assert sum(line.quantity == 0 for line in lines) == 568
    assert sum(line.quantity is not None and line.quantity > 0 for line in lines) == 48
    assert sum(line.quantity is not None and line.quantity < 0 for line in lines) == 0

    repeated = create_import_preview(
        db_session,
        "renamed-same-content.xlsx",
        content,
        business_batch_key="QINSI-5412-LOCAL-ACCEPTANCE",
    )
    assert repeated.id == batch.id
    assert repeated.status == "completed_with_issues"
    assert db_session.scalar(select(func.count()).select_from(QinsiImportBatch)) == 1
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1974
    assert db_session.scalar(select(func.count()).select_from(QinsiInventorySnapshot)) == 1


def test_background_preview_status_page_and_confirm_flow(client):
    http, db, _ = client
    response = http.post(
        "/products/import/preview",
        files={
            "file": (
                REFERENCE_BOOK.name,
                REFERENCE_BOOK.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"business_batch_key": "QINSI-PAGE-FLOW"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    batch = db.scalar(select(QinsiImportBatch))
    db.refresh(batch)
    assert batch.status == "previewed"
    status = http.get(f"/products/import/{batch.id}/status")
    page = http.get(f"/products/import?job_id={batch.id}")
    assert status.status_code == 200
    assert status.json()["new_count"] == 2
    assert page.status_code == 200
    assert "新增" in page.text and "配置 Master" in page.text

    confirmed = http.post(f"/products/import/{batch.id}/confirm", follow_redirects=False)
    assert confirmed.status_code == 303
    db.expire_all()
    batch = db.get(QinsiImportBatch, batch.id)
    assert batch.status == "completed"
    assert db.scalar(select(func.count()).select_from(Product)) == 2


FORMAL_HEADERS = [
    "商品名称", "商品规格", "货号", "商品条码", "单品条码", "型号规格", "图片",
    "品牌", "分类", "采购价", "销售价",
]


def formal_workbook_bytes(*, sheet_name: str, rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(FORMAL_HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in FORMAL_HEADERS])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def formal_row(**overrides) -> dict:
    row = {
        "商品名称": "秦丝正式商品",
        "商品规格": "30g",
        "货号": "QINSI-FORMAL-1",
        "商品条码": "",
        "单品条码": "4901234567894",
        "型号规格": "M-1",
        "图片": "https://images.qinsilk.com/formal.jpg",
        "品牌": "秦丝品牌",
        "分类": "秦丝分类",
        "采购价": "120",
        "销售价": "220",
    }
    row.update(overrides)
    return row


def test_formal_qinsi_product_list_sheet1_parses_by_headers():
    workbook = parse_qinsi_workbook(formal_workbook_bytes(sheet_name="Sheet1", rows=[formal_row()]))

    assert workbook.sheet_name == "Sheet1"
    assert workbook.source_format == "qinsi_product_list"
    assert len(workbook.rows) == 1
    mapped, warnings, errors = _mapped_row(workbook.rows[0][1], workbook.config)
    assert not warnings and not errors
    assert mapped["name_cn"] == "秦丝正式商品"
    assert mapped["qinsi_product_code"] == "QINSI-FORMAL-1"
    assert mapped["jan"] == "4901234567894"
    assert mapped["image_url"] == "https://images.qinsilk.com/formal.jpg"


def test_formal_qinsi_product_list_product_import_sheet_name_is_compatible():
    workbook = parse_qinsi_workbook(formal_workbook_bytes(sheet_name="商品导入", rows=[formal_row()]))

    assert workbook.sheet_name == "商品导入"
    assert workbook.source_format == "qinsi_product_list"


def test_formal_qinsi_product_list_arbitrary_sheet_name_is_detected_by_headers():
    workbook = parse_qinsi_workbook(formal_workbook_bytes(sheet_name="秦丝正式导出", rows=[formal_row()]))

    assert workbook.sheet_name == "秦丝正式导出"
    assert workbook.source_format == "qinsi_product_list"


def test_products_import_preview_uses_formal_parser_regardless_of_checkbox(client):
    http, db, _ = client
    content = formal_workbook_bytes(sheet_name="Sheet1", rows=[formal_row()])
    for checked in (False, True):
        response = http.post(
            "/products/import/preview",
            files={"file": ("formal.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"use_reference": REFERENCE_BOOK.name} if checked else {},
            follow_redirects=False,
        )
        assert response.status_code == 303
        batch = db.scalar(select(QinsiImportBatch).order_by(QinsiImportBatch.id.desc()))
        assert batch.status == "previewed"
        assert batch.error_message is None
        summary = json.loads(batch.summary_json)
        assert summary["sheet"] == "Sheet1"
        assert summary["source_format"] == "qinsi_product_list"


def test_formal_import_overwrites_qinsi_name_image_keeps_purchase_facts_and_unique_jan(db_session):
    location = Location(internal_code="LOC-QINSI-TEST", display_name="新日本仓库", location_type="qinsi_warehouse", is_qinsi_warehouse=True)
    product = Product(
        jan="4901234567894",
        name_cn="本地人工名",
        display_name="本地人工名",
        name_locked=True,
        product_data_confirmed=True,
        main_image_locked=True,
        main_image_path="data/products/main/manual.jpg",
        display_image_url="/product-images/1",
        main_image_source_url="https://example.test/old-auto.jpg",
        image_url="https://example.test/old-qinsi.jpg",
        status="new_pending_review",
    )
    db_session.add_all([location, product])
    db_session.flush()
    batch = ReceiptBatch(batch_no="FACT-QINSI", status="confirmed", image_status="ready", gpt_status="reviewed")
    receipt = Receipt(batch=batch, raw_store_name="采购事实店", confirmation_status="confirmed", review_status="reviewed")
    db_session.add_all([batch, receipt])
    db_session.flush()
    receipt_item = ReceiptItem(
        receipt=receipt,
        line_no=1,
        raw_name="小票原始名",
        jan_candidate=product.jan,
        product_id=product.id,
        match_status="matched_existing",
        quantity=3,
        unit_price=111,
        discount_amount=7,
        line_total=326,
        review_status="confirmed",
    )
    purchase = PurchaseBatch(
        batch_no="PB-FACT-QINSI",
        receipt=receipt,
        gpt_batch_id=batch.id,
        status="confirmed",
        confirmed_at=datetime.now(timezone.utc),
        default_initial_location_id=location.id,
        default_qinsi_warehouse_id=location.id,
    )
    purchase_item = PurchaseBatchItem(
        purchase_batch=purchase,
        product=product,
        receipt_item=receipt_item,
        quantity=3,
        unit_price=111,
        discount_amount=7,
        actual_line_amount=326,
        initial_location_id=location.id,
        qinsi_target_warehouse_id=location.id,
    )
    db_session.add_all([receipt_item, purchase, purchase_item])
    db_session.commit()

    content = formal_workbook_bytes(sheet_name="Sheet1", rows=[formal_row(
        **{"商品名称": "秦丝权威名", "图片": "https://images.qinsilk.com/authority.jpg"}
    )])
    preview = create_import_preview(db_session, "formal-authority.xlsx", content)
    assert preview.update_count == 1
    confirm_import(db_session, preview)
    db_session.refresh(product)
    db_session.refresh(receipt_item)
    db_session.refresh(purchase_item)

    assert product.name_cn == "秦丝权威名"
    assert product.name_ja is None
    # No name_ja means display_name is just the Chinese name -- no "|日文名待补"
    # placeholder tacked on for the missing side.
    assert product.display_name == "秦丝权威名"
    assert "本地人工名" not in product.display_name
    assert product.main_image_source_url == "https://images.qinsilk.com/authority.jpg"
    assert product.image_url == "https://images.qinsilk.com/authority.jpg"
    assert product.display_image_url == "https://images.qinsilk.com/authority.jpg"
    assert product.main_image_path is None
    assert product.main_image_locked is False
    assert receipt_item.raw_name == "小票原始名"
    assert (receipt_item.quantity, receipt_item.unit_price, receipt_item.discount_amount, receipt_item.line_total) == (3, 111, 7, 326)
    assert (purchase_item.quantity, purchase_item.unit_price, purchase_item.discount_amount, purchase_item.actual_line_amount) == (3, 111, 7, 326)
    assert db_session.scalar(
        select(func.count()).select_from(Product).where(Product.jan == "4901234567894")
    ) == 1


def test_qinsi_import_creates_derived_barcode_alias_idempotently(db_session):
    reproduction_jan = "4550726010198"
    workbook = load_workbook(BytesIO(REFERENCE_BOOK.read_bytes()))
    sheet = workbook["商品导入"]
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}
    code_column = headers.get("货号（必填且唯一）") or headers["货号(必填且唯一)"]
    barcode_column = headers["条码"]
    sheet.cell(2, code_column).value = f"/{reproduction_jan}"
    sheet.cell(2, barcode_column).value = None
    if sheet.max_row > 2:
        sheet.delete_rows(3, sheet.max_row - 2)
    output = BytesIO()
    workbook.save(output)

    batch = create_import_preview(
        db_session,
        "derived-qinsi-barcode.xlsx",
        output.getvalue(),
        business_batch_key="QINSI-DERIVED-BARCODE",
    )
    assert batch.status == "previewed"
    confirm_import(db_session, batch)
    product = db_session.scalar(
        select(Product).where(Product.qinsi_product_code == f"/{reproduction_jan}")
    )
    alias = db_session.scalar(
        select(ProductBarcode).where(ProductBarcode.barcode == reproduction_jan)
    )
    assert product is not None
    assert alias.product_id == product.id
    assert alias.source_system == "qinsi_sku_derived"
    assert alias.is_primary is False

    repeated = create_import_preview(
        db_session,
        "derived-qinsi-barcode-renamed.xlsx",
        output.getvalue(),
        business_batch_key="QINSI-DERIVED-BARCODE",
    )
    assert repeated.id == batch.id
    assert db_session.scalar(
        select(func.count()).select_from(ProductBarcode).where(
            ProductBarcode.barcode == reproduction_jan
        )
    ) == 1


def test_qinsi_import_maps_japanese_name_and_queues_changed_image_without_clearing_old_display(
    db_session,
):
    def workbook_bytes(image_url: str) -> bytes:
        workbook = load_workbook(BytesIO(REFERENCE_BOOK.read_bytes()))
        sheet = workbook["商品导入"]
        headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value}
        code_column = headers.get("货号（必填且唯一）") or headers["货号(必填且唯一)"]
        sheet.cell(2, code_column).value = "IMG-CHANGE-1"
        sheet.cell(2, headers["商品图片链接"]).value = image_url
        sheet.cell(2, headers["商品备注"]).value = "  日本語商品名  "
        if sheet.max_row > 2:
            sheet.delete_rows(3, sheet.max_row - 2)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    old_url = "https://images.qinsilk.com/old.png"
    new_url = "https://images.qinsilk.com/new.png"
    first = create_import_preview(db_session, "image-old.xlsx", workbook_bytes(old_url))
    confirm_import(db_session, first)
    product = db_session.scalar(select(Product).where(Product.qinsi_product_code == "IMG-CHANGE-1"))
    assert product.name_ja == "日本語商品名"
    assert product.product_note is None
    assert product.image_url == old_url
    product.display_image_url = f"/product-local-images/{product.id}?v=old"
    product.image_localization_source_url = old_url
    product.image_localization_status = "COMPLETED"
    db_session.commit()

    second = create_import_preview(db_session, "image-new.xlsx", workbook_bytes(new_url))
    assert second.update_count == 1
    confirm_import(db_session, second)
    db_session.refresh(product)
    jobs = list(db_session.scalars(
        select(DurableBackgroundJob)
        .where(DurableBackgroundJob.job_type == "PRODUCT_IMAGE_LOCALIZATION")
        .order_by(DurableBackgroundJob.id)
    ))
    assert len(jobs) == 2
    assert product.image_url == new_url
    assert product.main_image_source_url == new_url
    assert product.display_image_url == new_url
    assert product.local_image_path is None
    assert product.image_localization_source_url is None
    assert product.image_localization_status == "PENDING"
