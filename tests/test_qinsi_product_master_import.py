from __future__ import annotations

from io import BytesIO
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.field_purchase import create_field_batch, product_lookup_payload
from app.models import (
    FieldPurchaseItem,
    Location,
    Product,
    ProductBarcode,
    ProductOperationLog,
    PurchaseBatch,
    PurchaseBatchItem,
    QinsiConflictResolution,
    QinsiGoodsImportRow,
    QinsiImportBatch,
    QinsiProductMapping,
    Receipt,
    ReceiptBatch,
    ReceiptItem,
)
from app.qinsi_product_master_import import (
    MasterInputFile,
    confirm_qinsi_master_import,
    create_qinsi_master_preview,
    determine_qinsi_master_jan,
    export_qinsi_master_audit_workbook,
    resolve_qinsi_master_conflict,
)


HEADERS = [
    "商品名称", "商品规格", "货号", "商品条码", "单品条码", "型号规格", "图片链接",
    "品牌", "分类", "单位", "采购价", "销售价", "最低销售价", "保质期", "产地",
    "适用年龄", "排序", "状态", "库存预警下限", "库存预警上限", "备注",
]


def workbook_bytes(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row.get(header, "") for header in HEADERS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def master_file(rows: list[dict], filename: str = "qinsi-master.xlsx") -> MasterInputFile:
    return MasterInputFile(filename=filename, content=workbook_bytes(rows))


def qrow(*, name: str, goods_no: str, unit_barcode: str | None = None, product_barcode: str | None = None, spec: str | None = None) -> dict:
    row = {
        HEADERS[0]: name,
        HEADERS[2]: goods_no,
        HEADERS[17]: "enabled",
    }
    if spec is not None:
        row[HEADERS[1]] = spec
    if product_barcode is not None:
        row[HEADERS[3]] = product_barcode
    if unit_barcode is not None:
        row[HEADERS[4]] = unit_barcode
    return row


def test_qinsi_master_jan_priority_and_check_digit():
    assert determine_qinsi_master_jan(
        unit_barcode="4571609352419",
        product_barcode="4901234567894",
        goods_no="00123457",
    ) == ("4571609352419", "unit_barcode")
    assert determine_qinsi_master_jan(
        unit_barcode="BAD-UNIT",
        product_barcode="4901234567894",
        goods_no="00123457",
    ) == ("4901234567894", "product_barcode")
    assert determine_qinsi_master_jan(
        unit_barcode="4901234567895",
        product_barcode="not-a-gtin",
        goods_no="036000291452",
    ) == ("036000291452", "goods_no")
    assert determine_qinsi_master_jan(
        unit_barcode="4901234567895",
        product_barcode="12345678",
        goods_no="NOJAN-1",
    ) == (None, "none")


def test_qinsi_master_import_creates_no_jan_and_reimports_by_goods_no(db_session):
    batch = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "无JAN扭蛋",
        "货号": "NOJAN-1",
        "商品条码": "4901234567895",
        "单品条码": "",
        "图片链接": "https://example.test/nojan.jpg",
        "采购价": "120",
        "销售价": "220",
        "状态": "启用",
        "备注": "特殊商品",
    }])])
    assert batch.status == "previewed"
    assert batch.conflict_count == 0
    assert batch.error_count == 0
    assert batch.summary_json and '"no_jan_count": 1' in batch.summary_json

    confirm_qinsi_master_import(db_session, batch)
    product = db_session.scalar(select(Product).where(Product.qinsi_product_code == "NOJAN-1"))
    assert product is not None
    assert product.jan is None
    assert product.has_jan is False
    assert product.status == "qinsi_product_imported"
    assert product.source == "qinsi_import"
    assert product.qinsi_name == "无JAN扭蛋"

    second = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "无JAN扭蛋新版",
        "货号": "NOJAN-1",
        "状态": "启用",
    }], "qinsi-master-2.xlsx")])
    assert second.update_count == 1
    confirm_qinsi_master_import(db_session, second)
    db_session.refresh(product)
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1
    assert product.qinsi_name == "无JAN扭蛋新版"


def test_qinsi_master_import_existing_jan_updates_without_duplicate_and_overwrites_locked_name(db_session):
    existing = Product(
        jan="4901234567894",
        name_cn="高质量中文名",
        name_ja="正式日文名",
        display_name="高质量中文名|正式日文名",
        name_locked=True,
        product_data_confirmed=True,
        status="active",
    )
    db_session.add(existing)
    db_session.commit()

    batch = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "秦丝名称覆盖",
        "货号": "QINSI-JAN-1",
        "商品条码": "4901234567894",
        "品牌": "秦丝品牌",
        "分类": "秦丝分类",
        "单位": "个",
        "状态": "启用",
    }])])
    assert batch.update_count == 1
    confirm_qinsi_master_import(db_session, batch)
    db_session.refresh(existing)
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1
    assert existing.name_cn == "秦丝名称覆盖"
    assert existing.name_ja is None
    assert existing.display_name.startswith("秦丝名称覆盖|")
    assert "高质量中文名" not in existing.display_name
    assert existing.name_locked is False
    assert existing.product_data_confirmed is False
    assert existing.qinsi_product_code == "QINSI-JAN-1"
    assert existing.qinsi_product_imported if hasattr(existing, "qinsi_product_imported") else existing.status == "qinsi_product_imported"
    assert existing.qinsi_product_barcode == "4901234567894"
    assert existing.qinsi_brand == "秦丝品牌"


def test_qinsi_master_goods_no_gtin_updates_existing_goods_mapping(db_session):
    existing = Product(qinsi_product_code="036000291452", name_cn="既有无JAN秦丝商品", status="qinsi_product_imported")
    db_session.add(existing)
    db_session.commit()

    batch = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "货号是合法GTIN",
        "货号": "036000291452",
        "商品条码": "",
        "单品条码": "",
        "状态": "启用",
    }])])
    assert batch.update_count == 1
    assert batch.new_count == 0
    confirm_qinsi_master_import(db_session, batch)
    db_session.refresh(existing)
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1
    assert existing.jan == "036000291452"
    assert existing.qinsi_product_code == "036000291452"


def test_qinsi_master_goods_no_match_corrects_old_jan_instead_of_conflict(db_session):
    existing = Product(jan="036000291452", qinsi_product_code="036000291452", name_cn="既有JAN商品")
    db_session.add(existing)
    db_session.commit()

    batch = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "货号被占用",
        "货号": "036000291452",
        "单品条码": "4901234567894",
        "状态": "启用",
    }])])
    assert batch.conflict_count == 0
    assert batch.update_count == 1
    confirm_qinsi_master_import(db_session, batch)
    db_session.refresh(existing)
    assert existing.id == db_session.scalar(select(Product.id).where(Product.qinsi_product_code == "036000291452"))
    assert existing.jan == "4901234567894"


def test_qinsi_master_duplicate_valid_jan_in_excel_is_conflict(client):
    http, db, _ = client
    response = http.post(
        "/products/qinsi-master-import/preview",
        files=[
            ("files", ("master.xlsx", workbook_bytes([
                {"商品名称": "重复A", "货号": "71100129", "单品条码": "4571609352419", "状态": "启用"},
                {"商品名称": "重复B", "货号": "71100099", "商品条码": "4571609352419", "状态": "启用"},
            ]), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ],
        follow_redirects=False,
    )
    assert response.status_code == 303
    batch = db.scalar(select(QinsiImportBatch))
    assert batch.conflict_count == 2
    assert batch.new_count == 0
    assert '"same_jan_multi_group_count": 1' in batch.summary_json
    assert '"same_jan_multi_product_count": 2' in batch.summary_json
    rows = list(db.scalars(select(QinsiGoodsImportRow).order_by(QinsiGoodsImportRow.excel_row_number)))
    assert all(row.validation_status == "conflict" for row in rows)
    assert all("Excel内同一JAN出现多个秦丝商品" in (row.conflict_json or "") for row in rows)
    confirm = http.post(f"/products/qinsi-master-import/{batch.id}/confirm", follow_redirects=False)
    assert confirm.status_code == 303
    assert db.scalar(select(func.count()).select_from(Product).where(Product.jan == "4571609352419")) == 0
    assert http.get("/products/qinsi-master-import").status_code == 200
    assert http.get("/health").status_code == 200


def test_qinsi_shared_barcode_variant_can_be_confirmed_and_scans_as_candidates(db_session):
    shared = "4571609352419"
    batch = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="shared red", spec="red", goods_no="SHARED-RED", unit_barcode=shared),
        qrow(name="shared blue", spec="blue", goods_no="SHARED-BLUE", unit_barcode=shared),
    ])])
    assert batch.conflict_count == 2
    row = db_session.scalar(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.qinsi_product_code == "SHARED-RED"))

    resolution = resolve_qinsi_master_conflict(
        db_session,
        batch,
        row.id,
        resolution_type="shared_barcode_variant",
        note="different color",
        auto_apply=True,
    )
    assert resolution.conflict_key.endswith("|shared_barcode_variant")
    db_session.refresh(batch)
    assert batch.conflict_count == 0
    assert batch.new_count == 2

    confirm_qinsi_master_import(db_session, batch)
    products = list(db_session.scalars(select(Product).where(Product.qinsi_product_code.in_(["SHARED-RED", "SHARED-BLUE"])).order_by(Product.qinsi_product_code)))
    assert len(products) == 2
    assert {product.jan for product in products} == {None}
    assert db_session.scalar(select(func.count()).select_from(ProductBarcode).where(ProductBarcode.barcode == shared)) == 2
    lookup = product_lookup_payload(db_session, shared)
    assert lookup["status"] == "AMBIGUOUS"
    assert set(lookup["candidate_product_ids"]) == {product.id for product in products}


def test_qinsi_shared_barcode_resolution_auto_applies_without_excel_row_number(db_session):
    shared = "4571609352419"
    first = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="shared A", goods_no="AUTO-A", unit_barcode=shared),
        qrow(name="shared B", goods_no="AUTO-B", unit_barcode=shared),
    ], "auto-first.xlsx")])
    row = db_session.scalar(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == first.id, QinsiGoodsImportRow.qinsi_product_code == "AUTO-A"))
    resolve_qinsi_master_conflict(db_session, first, row.id, resolution_type="shared_barcode_variant", auto_apply=True)

    second = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="shared B", goods_no="AUTO-B", unit_barcode=shared),
        qrow(name="shared A", goods_no="AUTO-A", unit_barcode=shared),
    ], "auto-second.xlsx")])
    assert second.conflict_count == 0
    rows = list(db_session.scalars(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == second.id)))
    assert all("已按历史规则处理" in (row.warnings or "") for row in rows)


def test_qinsi_code_barcode_overlap_is_not_hard_conflict(db_session):
    batch = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="A", goods_no="CODE-A", unit_barcode="4571609352419"),
        qrow(name="B", goods_no="4571609352419", unit_barcode="4901234567894"),
    ])])
    assert batch.conflict_count == 0
    assert batch.new_count == 2


def test_qinsi_true_duplicate_requires_primary_product_choice(db_session):
    product = Product(
        jan="4571609352419",
        qinsi_product_code="DUP-PRIMARY",
        name_cn="primary",
        status="qinsi_product_imported",
        product_origin="qinsi",
    )
    db_session.add(product)
    db_session.commit()
    batch = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="duplicate", goods_no="DUP-OTHER", unit_barcode="4571609352419"),
    ])])
    row = db_session.scalar(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == batch.id))
    with pytest.raises(ValueError):
        resolve_qinsi_master_conflict(db_session, batch, row.id, resolution_type="true_duplicate")
    assert db_session.scalar(select(func.count()).select_from(QinsiConflictResolution)) == 0


def test_qinsi_shared_barcode_import_does_not_modify_purchase_facts(db_session):
    product = Product(
        qinsi_product_code="FACT-A",
        name_cn="fact product",
        status="qinsi_product_imported",
        product_origin="qinsi",
    )
    local_location = Location(internal_code="LOCAL-FACT", display_name="Local", location_type="local_physical")
    qinsi_location = Location(
        internal_code="QINSI-FACT",
        display_name="Qinsi",
        location_type="qinsi_warehouse",
        is_qinsi_warehouse=True,
    )
    db_session.add_all([product, local_location, qinsi_location])
    db_session.flush()
    receipt_batch = ReceiptBatch(batch_no="RB-FACT", request_id="rb-fact")
    db_session.add(receipt_batch)
    db_session.flush()
    receipt = Receipt(batch_id=receipt_batch.id, confirmation_status="confirmed")
    db_session.add(receipt)
    db_session.flush()
    receipt_item = ReceiptItem(
        receipt_id=receipt.id,
        line_no=1,
        raw_name="fact item",
        product_id=product.id,
        quantity=3,
        unit_price=120,
        line_total=360,
        confidence=1,
    )
    db_session.add(receipt_item)
    db_session.flush()
    purchase_batch = PurchaseBatch(
        batch_no="PB-FACT",
        receipt_id=receipt.id,
        gpt_batch_id=receipt_batch.id,
        confirmed_at=datetime.now(timezone.utc),
        default_initial_location_id=local_location.id,
        default_qinsi_warehouse_id=qinsi_location.id,
    )
    db_session.add(purchase_batch)
    db_session.flush()
    purchase_item = PurchaseBatchItem(
        purchase_batch_id=purchase_batch.id,
        product_id=product.id,
        receipt_item_id=receipt_item.id,
        quantity=3,
        unit_price=120,
        actual_line_amount=360,
        initial_location_id=local_location.id,
        qinsi_target_warehouse_id=qinsi_location.id,
    )
    db_session.add(purchase_item)
    db_session.commit()
    before = (receipt_item.product_id, receipt_item.quantity, receipt_item.unit_price, purchase_item.product_id, purchase_item.quantity, purchase_item.unit_price)

    shared = "4571609352419"
    batch = create_qinsi_master_preview(db_session, [master_file([
        qrow(name="fact A", goods_no="FACT-A", unit_barcode=shared),
        qrow(name="fact B", goods_no="FACT-B", unit_barcode=shared),
    ])])
    row = db_session.scalar(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.qinsi_product_code == "FACT-A"))
    resolve_qinsi_master_conflict(db_session, batch, row.id, resolution_type="shared_barcode_variant")
    confirm_qinsi_master_import(db_session, batch)
    db_session.refresh(receipt_item)
    db_session.refresh(purchase_item)
    after = (receipt_item.product_id, receipt_item.quantity, receipt_item.unit_price, purchase_item.product_id, purchase_item.quantity, purchase_item.unit_price)
    assert after == before


def test_qinsi_master_import_merges_local_jan_holder_into_existing_qinsi_product(db_session):
    qinsi = Product(qinsi_product_code="QINSI-MERGE", name_cn="秦丝旧商品", status="qinsi_product_imported", product_origin="qinsi")
    local = Product(jan="4571609352419", name_cn="本地临时商品", status="new_pending_review", product_origin="receipt")
    db_session.add_all([qinsi, local])
    db_session.flush()
    batch = create_field_batch(
        db_session,
        store_id=None,
        operator_name="采购员A",
        client_request_id="qinsi-local-jan-holder",
    )
    db_session.add(FieldPurchaseItem(
        batch_id=batch.id,
        product_id=local.id,
        jan=local.jan,
        quantity=2,
        status="CONFIRMED",
        captured_by="采购员A",
    ))
    db_session.commit()

    preview = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "秦丝最新商品",
        "货号": "QINSI-MERGE",
        "单品条码": "4571609352419",
        "状态": "启用",
    }])])
    assert preview.update_count == 1
    confirm_qinsi_master_import(db_session, preview)

    db_session.refresh(qinsi)
    assert qinsi.jan == "4571609352419"
    assert qinsi.qinsi_name == "秦丝最新商品"
    assert db_session.scalar(select(FieldPurchaseItem.product_id)) == qinsi.id
    assert db_session.get(Product, local.id) is None
    assert db_session.scalar(select(func.count()).select_from(Product).where(Product.jan == "4571609352419")) == 1


def test_qinsi_master_goods_no_updates_original_product_corrects_jan_and_keeps_local_history(db_session):
    product = Product(
        jan="4901234567894",
        qinsi_product_code="QINSI-KEEP",
        name_cn="本地小票名",
        name_source="receipt",
        local_image_path="data/uploads/local.jpg",
        main_image_path="data/uploads/main.jpg",
        display_image_url="/product-images/1",
        main_image_locked=True,
        image_localization_status="completed",
        status="qinsi_product_imported",
    )
    db_session.add(product)
    db_session.commit()
    batch = create_field_batch(
        db_session,
        store_id=None,
        operator_name="采购员A",
        client_request_id="qinsi-jan-correction-history",
    )
    db_session.add(FieldPurchaseItem(
        batch_id=batch.id,
        product_id=product.id,
        jan=product.jan,
        quantity=2,
        status="CONFIRMED",
        captured_by="采购员A",
    ))
    db_session.commit()

    preview = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "秦丝最新名",
        "货号": "QINSI-KEEP",
        "单品条码": "4571609352419",
        "图片链接": "https://example.test/qinsi.jpg",
        "品牌": "秦丝品牌",
        "分类": "秦丝分类",
        "单位": "盒",
        "采购价": "111.00",
        "销售价": "222.00",
        "状态": "启用",
        "备注": "最新备注",
    }])])
    assert preview.update_count == 1
    assert '"jan_correction_count": 1' in preview.summary_json
    confirm_qinsi_master_import(db_session, preview)
    db_session.refresh(product)
    assert product.id == db_session.scalar(select(FieldPurchaseItem.product_id))
    assert product.jan == "4571609352419"
    assert product.qinsi_name == "秦丝最新名"
    assert product.brand == "秦丝品牌"
    assert product.category == "秦丝分类"
    assert product.unit_name == "盒"
    assert product.purchase_price == 111
    assert product.sale_price == 222
    assert product.name_cn == "秦丝最新名"
    assert product.image_url == "https://example.test/qinsi.jpg"
    assert product.main_image_source_url == "https://example.test/qinsi.jpg"
    assert product.display_image_url == "https://example.test/qinsi.jpg"
    assert product.local_image_path is None
    assert product.main_image_path is None
    assert product.main_image_locked is False
    assert product.image_localization_status is None
    assert db_session.scalar(select(func.count()).select_from(Product)) == 1
    log = db_session.scalar(select(ProductOperationLog).where(ProductOperationLog.product_id == product.id))
    assert log is not None
    assert "4901234567894" in log.before_json
    assert "4571609352419" in log.after_json


def test_qinsi_master_same_goods_no_multiple_products_is_true_conflict(db_session):
    first = Product(qinsi_product_code="DUP-GOODS", name_cn="字段占用", status="qinsi_product_imported")
    second = Product(name_cn="映射占用", status="qinsi_product_imported")
    db_session.add_all([first, second])
    db_session.flush()
    db_session.add(QinsiProductMapping(qinsi_product_code="DUP-GOODS", product_id=second.id))
    db_session.commit()

    batch = create_qinsi_master_preview(db_session, [master_file([{
        "商品名称": "同货号冲突",
        "货号": "DUP-GOODS",
        "状态": "启用",
    }])])
    assert batch.conflict_count == 1
    row = db_session.scalar(select(QinsiGoodsImportRow).where(QinsiGoodsImportRow.import_batch_id == batch.id))
    assert "同一qinsi_goods_no对应多个本地Product" in row.conflict_json


def test_qinsi_master_audit_exports_97_conflicts_one_error_and_text_identifiers(client):
    http, db, _ = client
    local = Product(
        jan="4571609352419",
        qinsi_product_code="71100129",
        qinsi_product_barcode="P-LOCAL",
        qinsi_unit_barcode="U-LOCAL",
        name_cn="本地自嘲熊",
        status="qinsi_product_imported",
    )
    db.add(local)
    db.flush()
    batch = QinsiImportBatch(
        source_system="qinsi_product_master",
        original_filename="master.xlsx",
        file_hash="audit-test",
        file_content=b"{}",
        status="completed_with_issues",
        parse_version=4,
        total_rows=98,
        conflict_count=97,
        error_count=1,
    )
    db.add(batch)
    db.flush()
    for index in range(97):
        jan = "4571609352419" if index == 0 else f"0001234567{index:03d}"
        goods_no = "71100129" if index == 0 else f"0009876543{index:03d}"
        db.add(QinsiGoodsImportRow(
            import_batch_id=batch.id,
            source_file_name="master.xlsx",
            sheet_name="Sheet1",
            excel_row_number=index + 1,
            qinsi_product_code=goods_no,
            barcode=jan,
            parsed_data=(
                "{"
                f"\"source_file_name\":\"master.xlsx\",\"excel_row_number\":{index + 2},"
                f"\"qinsi_name\":\"冲突商品{index}\",\"qinsi_goods_no\":\"{goods_no}\","
                f"\"qinsi_product_barcode\":\"{goods_no}\",\"qinsi_unit_barcode\":\"{jan}\","
                f"\"jan\":\"{jan}\""
                "}"
            ),
            raw_json="{}",
            validation_status="conflict",
            conflict_json=(
                "[{\"field\":\"判定JAN\",\"excel_value\":\"4571609352419\","
                "\"existing_value\":{\"product_id\":%d},"
                "\"message\":\"秦丝多个商品命中同一有效JAN：4571609352419\"}]"
            ) % local.id,
        ))
    db.add(QinsiGoodsImportRow(
        import_batch_id=batch.id,
        source_file_name="master.xlsx",
        sheet_name="Sheet1",
        excel_row_number=98,
        qinsi_product_code="ERR-1",
        barcode=None,
        parsed_data=(
            "{\"source_file_name\":\"master.xlsx\",\"excel_row_number\":99,"
            "\"qinsi_name\":\"错误商品\",\"qinsi_goods_no\":\"ERR-1\","
            "\"qinsi_product_barcode\":null,\"qinsi_unit_barcode\":null,\"jan\":null}"
        ),
        raw_json="{\"货号\":\"ERR-1\"}",
        validation_status="error",
        errors="[\"商品名称为空\"]",
    ))
    db.commit()

    content = export_qinsi_master_audit_workbook(db, batch)
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["冲突明细", "错误明细", "汇总"]
    conflict_sheet = workbook["冲突明细"]
    error_sheet = workbook["错误明细"]
    summary_sheet = workbook["汇总"]
    assert conflict_sheet.max_row == 98
    assert error_sheet.max_row == 2
    assert any(cell.value == "4571609352419" for row in conflict_sheet.iter_rows() for cell in row)
    assert conflict_sheet["D2"].data_type == "s"
    assert conflict_sheet["D2"].number_format == "@"
    assert conflict_sheet["G2"].data_type == "s"
    assert dict(summary_sheet.iter_rows(min_row=2, values_only=True))["冲突总数"] == 97
    assert dict(summary_sheet.iter_rows(min_row=2, values_only=True))["错误总数"] == 1
    assert http.get("/health").status_code == 200
